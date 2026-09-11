"""Loopback-only desktop workspace. The existing collector remains the parser."""
from __future__ import annotations

import argparse
import atexit
import copy
import json
import hashlib
import mimetypes
import booking_file_lock
import os
import re
import secrets
import shutil
import sys
import threading
import time
import traceback
import webbrowser
from dataclasses import asdict
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from urllib.request import urlopen, Request

SUPPORT = Path(__file__).resolve().parent
sys.path.insert(0, str(SUPPORT / "python_packages"))
if (SUPPORT / "ms-playwright").exists():
    os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(SUPPORT / "ms-playwright"))

import naver_room_booking_upgraded as core
from booking_store import Store, atomic_json, read_json, entry_key, fingerprint, exact_price, now
from booking_discovery import search_places, verify_place, place_id
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from room_classifier import PARSER_VERSION, filtered_snapshot, partition, item_key

VERSION = "2.1.0"


class Manager:
    def __init__(self, folder):
        self.store = Store(folder)
        self.lock = threading.RLock()
        self.stop = threading.Event()
        self.worker = None
        self.state_path = self.store.home / "workspace.json"
        self.data = read_json(self.state_path, {"results": {}, "proposals": {}, "discoveries": [], "jobs": []})
        self.data.setdefault("catalogs", {})
        self.legacy = self.store.legacy_results()
        self.legacy_stamp = time.monotonic()
        self.job = {"state": "idle", "message": "대기 중", "completed": 0, "total": 0, "logs": []}
        for job in self.data["jobs"]:
            if job["state"] == "running":
                job.update(state="interrupted", message="프로그램 종료로 작업이 중단되었습니다.", finished_at=now())
        self.persist()

    def persist(self):
        with self.lock:
            atomic_json(self.state_path, self.data)

    def busy(self):
        return bool(self.worker and self.worker.is_alive())

    def log(self, message):
        with self.lock:
            self.job["message"] = str(message)
            self.job["logs"] = (self.job["logs"] + [{"time": now(), "text": str(message)}])[-500:]

    def state(self):
        with self.lock:
            if not self.busy() and time.monotonic() - self.legacy_stamp > 30:
                self.legacy = self.store.legacy_results()
                self.legacy_stamp = time.monotonic()
            entries, rev = self.store.entries()
            archived = set(self.store.meta["archived"])
            rows = []
            for e in entries:
                key = entry_key(e)
                candidates = [r for r in (self.data["results"].get(key), self.legacy.get(key)) if r]
                result = max(candidates, key=lambda r: r.get("checked_at", "")) if candidates else None
                result = copy.deepcopy(result)
                if result:
                    result["stale"] = (result.get("fingerprint") != fingerprint(e)
                        or result.get("parser_version") != PARSER_VERSION
                        or result.get("rules_revision") != self.rules_revision(e.company_id))
                catalog = self.data["catalogs"].get(e.company_id)
                groups = partition(catalog["snapshot"], self.rules(e.company_id)) if catalog else {}
                rows.append({**asdict(e), "key": key, "region": self.store.meta["regions"].get(key, ""),
                    "archived": key in archived, "changed": key in self.store.meta["changed"],
                    "result": result, "proposal": self.data["proposals"].get(key), "price_amount": exact_price(e.price),
                    "item_groups": groups, "rules_revision": self.rules_revision(e.company_id),
                    "items_checked_at": catalog.get("checked_at") if catalog else None})
            return copy.deepcopy({"version": VERSION, "parser_version": PARSER_VERSION, "folder": str(self.store.folder), "today": date.today().isoformat(),
                "revision": rev, "entries": rows, "job": self.job, "jobs": self.data["jobs"][:30],
                "discoveries": self.data["discoveries"], "busy": self.busy()})

    def select(self, payload):
        entries, rev = self.store.entries()
        if payload.get("revision") != rev:
            raise ValueError("DB가 바뀌었습니다. 새로고침 후 다시 선택해 주세요.")
        active = [e for e in entries if entry_key(e) not in self.store.meta["archived"]]
        keys = payload.get("keys")
        if payload.get("all") is True:
            selected = active
        elif isinstance(keys, list) and keys:
            selected = [e for e in active if entry_key(e) in keys]
            if len(selected) != len(set(keys)):
                raise ValueError("선택 항목이 변경되었거나 보관되어 있습니다.")
        else:
            raise ValueError("작업할 항목을 선택해 주세요.")
        if not selected:
            raise ValueError("작업할 활성 항목이 없습니다.")
        return selected

    def start(self, kind, payload):
        with self.lock:
            if self.busy():
                raise ValueError("작업 중입니다. 완료 또는 중지 후 다시 실행해 주세요.")
            checkin = core.parse_date(payload.get("date") or date.today().isoformat())
            selected = self.select(payload) if kind in {"collect", "scan"} else []
            if kind not in {"collect", "scan", "discover", "verify"}:
                raise ValueError("알 수 없는 작업입니다.")
            if kind == "verify" and not payload.get("links", "").strip():
                raise ValueError("네이버 플레이스 링크 또는 업체번호를 입력해 주세요.")
            self.stop.clear()
            self.job = {"id": datetime.now().strftime("%Y%m%d_%H%M%S_%f"), "kind": kind,
                "state": "running", "date": checkin.isoformat(), "started_at": now(),
                "message": "작업 준비 중", "total": len(selected), "completed": 0, "normal": 0, "issues": 0, "logs": []}
            self.data["jobs"].insert(0, self.job)
            self.data["jobs"] = self.data["jobs"][:100]
            self.persist()
            self.worker = threading.Thread(target=self._execute, args=(kind, selected, checkin, payload), daemon=True)
            self.worker.start()
            return {"job_id": self.job["id"]}

    def _execute(self, kind, selected, checkin, payload):
        try:
            with core.collection_run_lock(self.store.folder):
                if kind in {"collect", "scan"}:
                    self._inspect(kind, selected, checkin)
                else:
                    self._discover(kind, checkin, payload)
            with self.lock:
                self.job["state"] = "stopped" if self.stop.is_set() else ("issues" if self.job["issues"] else "completed")
                self.log(f"{'중지' if self.stop.is_set() else '완료'}: 처리 {self.job['completed']}/{self.job['total']}, 정상 {self.job['normal']}, 확인 필요 {self.job['issues']}")
        except Exception as exc:
            with self.lock:
                self.job["state"] = "failed"
                self.log(str(exc))
                (self.store.home / (self.job["id"] + "_error.txt")).write_text(traceback.format_exc(), encoding="utf-8")
        finally:
            with self.lock:
                self.job["finished_at"] = now()
                self.persist()
                atomic_json(self.store.home / (self.job["id"] + ".json"), self.job)

    def _proposal(self, entry, snapshot, revision):
        snapshot = {**snapshot, "roomRules": self.rules(entry.company_id)}
        filtered, categories = filtered_snapshot(snapshot, self.rules(entry.company_id))
        detected = core.extract_detected_rooms(filtered, core.normalize_words(core.DEFAULT_STATUS_WORDS))
        old_keys = {core.normalize_name(r) for r in entry.rooms}
        new_keys = {core.normalize_name(r) for r in detected}
        group = [e for e in self.store.entries()[0] if e.company_id == entry.company_id]
        other_keys = {core.normalize_name(r) for e in group if e.excel_row != entry.excel_row for r in e.rooms}
        added = [r for r in detected if core.normalize_name(r) not in old_keys | other_keys]
        removed = [r for r in entry.rooms if core.normalize_name(r) not in new_keys]
        return {"key": entry_key(entry), "revision": revision, "checked_at": now(), "detected": detected,
            "added": added, "removed": removed, "split": len(group) > 1, "before": entry.rooms,
            "proposed": detected if len(group) == 1 and detected else entry.rooms,
            "changed": bool(added or removed or categories["review"]), "safe": bool(detected) and len(group) == 1 and not categories["review"],
            "review": [{"title": c["title"], "reason": c["reason"]} for c in categories["review"]],
            "excluded": [{"title": c["title"], "reason": c["reason"]} for c in categories["excluded"]],
            "parser_version": PARSER_VERSION}

    def _inspect(self, kind, entries, checkin):
        words = core.normalize_words(core.DEFAULT_STATUS_WORDS)
        _, rev = self.store.entries()
        wb = None
        collected_normal = []
        results = []
        if kind == "collect":
            if self.store.result.exists():
                shutil.copy2(self.store.result, self.store.home / ("현황_수집전_" + self.job["id"] + ".xlsx"))
                wb = load_workbook(self.store.result)
            else:
                wb = Workbook()
                wb.active.title = core.STATUS_SHEET_NAME
        try:
            with core.BrowserCollector(checkin, words, False, 25, self.log) as collector:
                for e in entries:
                    if self.stop.is_set():
                        break
                    key = entry_key(e)
                    try:
                        url, snapshot = collector.get_snapshot(e.company_id)
                        snapshot = {**snapshot, "roomRules": self.rules(e.company_id)}
                        with self.lock:
                            self.data["catalogs"][e.company_id] = {"snapshot": {"roomCards": snapshot.get("roomCards", []), "strictRoomCards": True}, "checked_at": now()}
                        proposal = self._proposal(e, snapshot, rev)
                        result = core.analyze_entry_snapshot(e, snapshot, words, "matched", url)
                        if result.status == "정상" and proposal["added"]:
                            result.status = f"신규 객실 {len(proposal['added'])}개 확인 필요"
                        if result.status == "정상" and proposal["review"]:
                            result.status = f"상품 판별 필요 {len(proposal['review'])}개: 객실/제외 선택 필요"
                        group = [v for v in self.store.entries()[0] if v.company_id == e.company_id and v.excel_row != e.excel_row]
                        if any(set(core.normalized_room_list(e.rooms)) & set(core.normalized_room_list(v.rooms)) for v in group):
                            result.status = "분류 간 객실 중복: 가동률 중복 집계 주의"
                        with self.lock:
                            self.data["proposals"][key] = proposal
                            if kind == "collect":
                                self.data["results"][key] = {"status": result.status, "reserved": result.reserved_count,
                                    "matched": result.matched_count, "total": result.saved_total,
                                    "db_count": result.db_room_count, "missing": result.missing_count,
                                    "reserved_rooms": result.reserved_rooms, "missing_rooms": result.missing_rooms,
                                    "checked_at": now(), "date": checkin.isoformat(), "fingerprint": fingerprint(e),
                                    "parser_version": PARSER_VERSION, "rules_revision": self.rules_revision(e.company_id),
                                    "excluded_items": proposal["excluded"], "review_items": proposal["review"]}
                            self.job["completed"] += 1
                            self.job["normal" if result.status == "정상" else "issues"] += 1
                        if wb is not None and result.status == "정상":
                            core.write_result_row(wb, result, checkin)
                            collected_normal.append(key)
                            ws = wb[e.sheet_title]
                            for row in range(2, ws.max_row + 1):
                                if core.same_day(ws.cell(row, 1).value, checkin):
                                    amount = exact_price(e.price)
                                    ws.cell(row, 6).value = result.reserved_count * amount if amount is not None else None
                        results.append(result)
                        self.log(f"{self.job['completed']}/{len(entries)} {e.major} / {e.sheet_title}: 마감 {result.reserved_count}/{result.saved_total}, 매칭 {result.matched_count}/{len(e.rooms)} · {result.status}")
                    except Exception as exc:
                        with self.lock:
                            self.job["completed"] += 1
                            self.job["issues"] += 1
                            self.data["proposals"][key] = {"key": key, "revision": rev, "error": str(exc), "changed": True,
                                "proposed": e.rooms, "before": e.rooms, "split": False, "safe": False, "added": [], "removed": [], "checked_at": now()}
                            if kind == "collect":
                                self.data["results"][key] = {"status": "오류: " + str(exc), "reserved": None, "total": None,
                                    "matched": 0, "missing": len(e.rooms), "db_count": len(e.rooms),
                                    "checked_at": now(), "date": checkin.isoformat(), "fingerprint": fingerprint(e),
                                    "parser_version": PARSER_VERSION, "rules_revision": self.rules_revision(e.company_id)}
                        self.log(e.sheet_title + ": " + str(exc))
                        if "접근 제한" in str(exc) or "접근을 제한" in str(exc):
                            raise
                    finally:
                        self.persist()
        finally:
            try:
                if wb is not None and collected_normal:
                    self._overview(wb, checkin)
                    temp = self.store.home / "collection_result.xlsx"
                    wb.save(temp)
                    os.replace(temp, self.store.result)
                    self.job["result_saved"] = True
                    self.store.meta["changed"] = [k for k in self.store.meta["changed"] if k not in collected_normal]
                    self.log("예약현황_Result.xlsx 저장 완료. 정상 매칭 항목만 반영했습니다.")
                if results and kind == "collect":
                    core.write_run_log(self.store.backup / ("booking_run_" + self.job["id"] + ".csv"), results)
                self.store.save_meta()
            finally:
                if wb:
                    wb.close()

    def _overview(self, wb, checkin):
        title = core.STATUS_SHEET_NAME
        ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title, 0)
        for merged in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merged))
        ws.delete_rows(1, ws.max_row)
        ws.append(["업체번호", "대분류", "중분류", "체크인", "마감 객실", "매칭 객실", "DB 객실", "마감률", "입력 금액", "추정 매출", "수집 상태", "확인 시각"])
        for entry in self.store.entries()[0]:
            key = entry_key(entry)
            if key in self.store.meta["archived"]:
                continue
            result = self.data["results"].get(key, {})
            status = result.get("status", "미수집")
            current = result.get("date") == checkin.isoformat() and result.get("fingerprint") == fingerprint(entry)
            if result and not current:
                status = "이전 기록 / 재수집 필요"
            valid = current and status == "정상"
            amount = exact_price(entry.price)
            total = result.get("total")
            reserved = result.get("reserved")
            ws.append([entry.company_id, entry.major, entry.sheet_title, result.get("date", ""), reserved,
                result.get("matched"), len(entry.rooms), reserved / total if valid and total else None,
                amount if amount is not None else entry.price, reserved * amount if valid and amount is not None else None,
                status, result.get("checked_at", "")])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="13734D")
            cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True, size=10)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=10)
                cell.alignment = Alignment(vertical="center")
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F1F6F3")
            row[7].number_format = "0.0%"
            row[8].number_format = "#,##0"
            row[9].number_format = "#,##0"
        for col, width in zip("ABCDEFGHIJKL", [18, 22, 30, 15, 12, 12, 12, 12, 18, 18, 42, 23]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

    def _discover(self, kind, checkin, payload):
        if kind == "discover":
            if not str(payload.get("region", "")).strip():
                raise ValueError("지역을 입력해 주세요.")
            query = str(payload.get("region", "")).strip() + " " + str(payload.get("keyword", "캠핑장"))
            candidates = search_places(query.strip(), min(20, max(1, int(payload.get("limit", 10)))), self.log, self.stop)
        else:
            ids = list(dict.fromkeys(place_id(v) for v in str(payload["links"]).splitlines() if v.strip()))
            if "" in ids or not ids or len(ids) > 20:
                raise ValueError("네이버 플레이스 링크 또는 업체번호를 한 줄에 하나씩 입력해 주세요. 최대 20개입니다.")
            candidates = [{"company_id": c, "name": "", "region": str(payload.get("region", ""))} for c in ids]
        with self.lock:
            self.data["discoveries"] = []
            self.job["total"] = len(candidates)
        with core.BrowserCollector(checkin, core.normalize_words(core.DEFAULT_STATUS_WORDS), False, 25, self.log) as collector:
            for candidate in candidates:
                if self.stop.is_set():
                    break
                verified = verify_place(collector, candidate)
                verified["existing"] = any(e.company_id == candidate["company_id"] for e in self.store.entries()[0])
                with self.lock:
                    self.data["discoveries"].append(verified)
                    self.job["completed"] += 1
                    self.job["normal" if verified["verified"] else "issues"] += 1
                self.log(verified["name"] + ": " + verified["status"])
                self.persist()

    def apply(self, payload):
        with self.lock:
            if self.busy():
                raise ValueError("작업 중에는 DB를 저장할 수 없습니다. 중지 후 저장해 주세요.")
            if payload.get("from_proposal"):
                for edit in payload.get("edits", []):
                    proposal = self.data["proposals"].get(edit.get("key"), {})
                    if proposal.get("parser_version") != PARSER_VERSION or proposal.get("review"):
                        raise ValueError("상품 판별에서 확인 필요 항목을 먼저 객실 또는 제외로 지정해 주세요. 이전 버전의 후보는 다시 점검해야 합니다.")
            result = self.store.apply(payload)
            for edit in payload["edits"]:
                self.data["proposals"].pop(edit.get("key"), None)
            for proposal in self.data["proposals"].values():
                if proposal["revision"] == payload["revision"]:
                    proposal["revision"] = result["revision"]
            self.persist()
            return result

    def rules(self, company):
        return self.store.meta.get("room_rules", {}).get(company, {})

    def rules_revision(self, company):
        return hashlib.sha256(json.dumps(self.rules(company), sort_keys=True, ensure_ascii=False).encode()).hexdigest()

    def set_rules(self, payload):
        with self.lock:
            if self.busy():
                raise ValueError("작업 완료 후 상품 판별 기준을 저장해 주세요.")
            with core.collection_run_lock(self.store.folder):
                entries, rev = self.store.entries()
                entry = next((e for e in entries if entry_key(e) == payload.get("key")), None)
                if not entry or payload.get("revision") != rev:
                    raise ValueError("DB가 변경되었습니다. 화면을 새로고침해 주세요.")
                company = entry.company_id
                if payload.get("rules_revision") != self.rules_revision(company):
                    raise ValueError("상품 판별 기준이 변경되었습니다. 다시 열어 주세요.")
                catalog = self.data["catalogs"].get(company)
                if not catalog:
                    raise ValueError("선택 업체의 객실명 점검을 먼저 실행해 주세요.")
                cards = {item_key(c): c for c in catalog["snapshot"]["roomCards"]}
                rules = copy.deepcopy(self.rules(company))
                for change in payload.get("items", []):
                    key, mode = change.get("itemKey"), change.get("mode")
                    if key not in cards or mode not in {"auto", "include", "exclude"}:
                        raise ValueError("확인할 수 없는 상품 판별 요청입니다.")
                    if mode == "auto":
                        rules.pop(key, None)
                    else:
                        rules[key] = {"mode": mode, "title": cards[key]["title"], "updated_at": now()}
                before = copy.deepcopy(self.store.meta)
                backup = self.store.home / ("상품판별_수정전_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f") + ".json")
                atomic_json(backup, before)
                self.store.meta.setdefault("room_rules", {})[company] = rules
                group = [e for e in entries if e.company_id == company]
                self.store.meta["changed"] = sorted(set(self.store.meta["changed"] + [entry_key(e) for e in group]))
                try:
                    self.store.save_meta()
                except Exception:
                    self.store.meta = before
                    raise
                for e in group:
                    self.data["proposals"][entry_key(e)] = self._proposal(e, catalog["snapshot"], rev)
                self.persist()
                return {"keys": [entry_key(e) for e in group], "backup": str(backup)}

    def archive(self, payload):
        with self.lock:
            if self.busy():
                raise ValueError("작업 완료 후 보관 상태를 변경해 주세요.")
            with core.collection_run_lock(self.store.folder):
                entries, rev = self.store.entries()
                if payload.get("revision") != rev:
                    raise ValueError("DB가 바뀌었습니다. 새로고침해 주세요.")
                keys = set(payload.get("keys", []))
                if not keys or not keys <= {entry_key(e) for e in entries}:
                    raise ValueError("보관할 항목을 선택해 주세요.")
                archived = set(self.store.meta["archived"])
                self.store.meta["archived"] = sorted(archived | keys if payload.get("archived", True) else archived - keys)
                self.store.save_meta()
                return {"count": len(keys)}


class Server(ThreadingHTTPServer):
    daemon_threads = True


def make_server(manager, port):
    token = secrets.token_urlsafe(32)

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def send(self, status, body, content_type="application/json; charset=utf-8"):
            data = json.dumps(body, ensure_ascii=False, default=str).encode("utf-8") if not isinstance(body, bytes) else body
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'")
            self.end_headers()
            self.wfile.write(data)

        def valid_host(self):
            return self.headers.get("Host") in {f"127.0.0.1:{self.server.server_port}", f"localhost:{self.server.server_port}"}

        def do_GET(self):
            if not self.valid_host():
                return self.send(403, {"error": "허용되지 않은 호스트"})
            path = urlparse(self.path).path
            try:
                if path == "/api/health":
                    return self.send(200, {"version": VERSION, "folder": str(manager.store.folder)})
                if path == "/api/state":
                    return self.send(200, manager.state())
                if path == "/api/history":
                    key = parse_qs(urlparse(self.path).query).get("key", [""])[0]
                    with manager.lock:
                        return self.send(200, {"history": manager.store.history(key)})
                if path in {"/download/db", "/download/result"}:
                    file = manager.store.db if path.endswith("db") else manager.store.result
                    return self.send(200, file.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                if path in {"/", "/app.js", "/style.css", "/lucide.min.js", "/favicon.svg"}:
                    file = SUPPORT / "web" / ("index.html" if path == "/" else path[1:])
                    body = file.read_bytes()
                    if path == "/":
                        body = body.replace(b"__TOKEN__", token.encode())
                    return self.send(200, body, (mimetypes.guess_type(str(file))[0] or "text/plain") + "; charset=utf-8")
                return self.send(404, {"error": "페이지 없음"})
            except Exception as exc:
                return self.send(500, {"error": str(exc)})

        def do_POST(self):
            origin = self.headers.get("Origin")
            valid_origin = not origin or origin in {f"http://127.0.0.1:{self.server.server_port}", f"http://localhost:{self.server.server_port}"}
            if not self.valid_host() or not valid_origin or self.headers.get("X-Booking-Token") != token:
                return self.send(403, {"error": "이 프로그램 화면에서 다시 요청해 주세요."})
            try:
                size = int(self.headers.get("Content-Length", 0))
                if size < 0 or size > 2_000_000:
                    raise ValueError("요청이 너무 큽니다.")
                data = json.loads(self.rfile.read(size) or b"{}")
                path = urlparse(self.path).path
                if path == "/api/start":
                    result = manager.start(data.get("kind"), data)
                elif path == "/api/stop":
                    manager.stop.set()
                    result = {"message": "현재 항목 처리 후 중지합니다."}
                elif path == "/api/apply":
                    result = manager.apply(data)
                elif path == "/api/archive":
                    result = manager.archive(data)
                elif path == "/api/room-rules":
                    result = manager.set_rules(data)
                else:
                    return self.send(404, {"error": "알 수 없는 요청"})
                self.send(200, result)
            except (ValueError, RuntimeError, PermissionError) as exc:
                self.send(409, {"error": str(exc)})
            except Exception as exc:
                self.send(500, {"error": str(exc)})

    return Server(("127.0.0.1", port), Handler)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--folder", type=Path, default=SUPPORT.parent)
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--no-open", action="store_true")
    parser.add_argument("--collect", action="store_true")
    parser.add_argument("--date", default=date.today().isoformat())
    args = parser.parse_args()
    args.folder = args.folder.resolve()
    for port in range(args.port, args.port + 20):
        address = f"http://127.0.0.1:{port}"
        try:
            with urlopen(address + "/api/health", timeout=0.2) as response:
                health = json.load(response)
            if health.get("folder") != str(args.folder) or health.get("version") != VERSION:
                continue
        except Exception:
            continue
        if args.collect:
            with urlopen(address) as response:
                token = re.search(r'name="booking-token" content="([^"]+)"', response.read().decode()).group(1)
            with urlopen(address + "/api/state") as response:
                state = json.load(response)
            request = Request(address + "/api/start", data=json.dumps({"kind": "collect", "date": args.date,
                "revision": state["revision"], "all": True}).encode(), headers={"Content-Type": "application/json", "X-Booking-Token": token})
            with urlopen(request) as response:
                job_id = json.load(response)["job_id"]
            while True:
                with urlopen(address + "/api/state") as response:
                    state = json.load(response)
                job = next((j for j in state["jobs"] if j["id"] == job_id), None)
                if job and job["state"] != "running":
                    print(json.dumps(job, ensure_ascii=False))
                    return 0 if job["state"] in {"completed", "issues"} else 1
                time.sleep(2)
        if not args.no_open:
            webbrowser.open(address)
        return 0
    service_home = args.folder / core.BACKUP_FOLDERNAME / "통합관리"
    service_home.mkdir(parents=True, exist_ok=True)
    guard = open(service_home / "manager_service.lock", "a+b")
    if guard.seek(0, os.SEEK_END) == 0:
        guard.write(b"0")
        guard.flush()
    guard.seek(0)
    try:
        booking_file_lock.acquire(guard)
    except OSError as exc:
        guard.close()
        raise RuntimeError("프로그램이 시작 중이거나 명령 수집이 실행 중입니다. 잠시 후 다시 열어 주세요.") from exc
    atexit.register(guard.close)
    manager = Manager(args.folder)
    if args.collect:
        manager.start("collect", {"date": args.date, "revision": manager.store.entries()[1], "all": True})
        manager.worker.join()
        print(json.dumps(manager.job, ensure_ascii=False))
        return 0 if manager.job["state"] in {"completed", "issues"} else 1
    for port in range(args.port, args.port + 20):
        try:
            server = make_server(manager, port)
            break
        except OSError:
            continue
    else:
        raise RuntimeError("사용 가능한 프로그램 포트를 찾지 못했습니다.")
    address = f"http://127.0.0.1:{server.server_port}"
    atomic_json(manager.store.home / "server.json", {"url": address, "pid": os.getpid(), "version": VERSION})
    if not args.no_open:
        webbrowser.open(address)
    print(address, flush=True)
    server.serve_forever()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        log_path = SUPPORT / "manager_start_error.txt"
        log_path.write_text(traceback.format_exc(), encoding="utf-8")
        if sys.stdout is None:
            import ctypes
            ctypes.windll.user32.MessageBoxW(0, f"시작 실패. 오류 기록: {log_path}", "캠핑 인사이트", 16)
        else:
            traceback.print_exc()
        raise SystemExit(1)
