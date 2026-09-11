"""Excel-backed catalog with revision checks and recoverable transactions."""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
from urllib.parse import parse_qs, urlparse
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
import naver_room_booking_upgraded as core


def now():
    return datetime.now().isoformat(timespec="seconds")


def revision(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest() if Path(path).exists() else ""


def atomic_json(path, value):
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temp, path)


def read_json(path, default):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def entry_key(entry):
    return hashlib.sha256(f"{entry.company_id}|{entry.sheet_title}".encode()).hexdigest()[:20]


def fingerprint(entry):
    value = [entry.company_id, entry.major, entry.sheet_title, entry.price, entry.rooms]
    return hashlib.sha256(json.dumps(value, ensure_ascii=False).encode()).hexdigest()


def clean_rooms(value):
    if isinstance(value, str):
        value = value.splitlines()
    if not isinstance(value, list) or len(value) > 2000:
        raise ValueError("객실 목록을 확인해 주세요. 최대 2,000개입니다.")
    rooms, keys = [], set()
    for item in value:
        if not isinstance(item, str):
            raise ValueError("객실명은 문자로 입력해 주세요.")
        item = item.strip()
        if not item:
            continue
        key = core.normalize_name(item)
        if key in keys:
            raise ValueError(f"중복 객실명: {item}")
        if len(item) > 180:
            raise ValueError("객실명이 너무 깁니다.")
        keys.add(key)
        rooms.append(item)
    return rooms


def exact_price(text):
    text = str(text or "").replace(",", "").replace(" ", "")
    if re.fullmatch(r"\d+(?:\.\d+)?만(?:원)?", text):
        return round(float(re.sub(r"만원?$", "", text)) * 10000)
    if re.fullmatch(r"\d+(?:\.0+)?원?", text):
        amount = int(float(text.rstrip("원")))
        return amount if amount >= 1000 or text.endswith("원") else None
    return None


def validate_entry(data, row):
    company = str(data.get("company_id", "")).strip()
    if not re.fullmatch(r"\d{5,15}", company):
        raise ValueError("업체번호는 네이버 플레이스 고유번호(숫자)로 입력해 주세요.")
    title = str(data.get("sheet_title", "")).strip()
    if not title or len(title) > 31 or re.search(r"[\\/*?:\[\]]", title) or title == core.STATUS_SHEET_NAME:
        raise ValueError("중분류는 31자 이내, 엑셀 시트에 사용할 수 있는 고유 이름이어야 합니다.")
    return core.DbEntry(row, company, str(data.get("major", "")).strip()[:120], title,
                        str(data.get("price", "")).strip()[:80], clean_rooms(data.get("rooms", [])))


class Store:
    def __init__(self, folder):
        self.folder = Path(folder)
        self.db = self.folder / core.DB_FILENAME
        self.result = self.folder / core.RESULT_FILENAME
        self.backup = self.folder / core.BACKUP_FOLDERNAME
        self.backup.mkdir(exist_ok=True)
        self.home = self.backup / "통합관리"
        self.home.mkdir(exist_ok=True)
        self.meta_path = self.home / "catalog.json"
        self.meta = read_json(self.meta_path, {"archived": [], "regions": {}, "changed": []})
        self.cache = None
        with core.collection_run_lock(self.folder):
            self.recover()

    def recover(self):
        journal = self.home / "transaction.json"
        if journal.exists():
            record = read_json(journal, {})
            # A crash between Excel replacements is rolled back before any new writes.
            for name in (core.DB_FILENAME, core.RESULT_FILENAME, "catalog.json"):
                backup = Path(record["backup"]) / name
                if backup.exists():
                    shutil.copy2(backup, self.meta_path if name == "catalog.json" else self.folder / name)
            journal.unlink()
            self.meta = read_json(self.meta_path, self.meta)

    def entries(self):
        stamp = (self.db.stat().st_mtime_ns, self.db.stat().st_size)
        if not self.cache or self.cache[0] != stamp:
            wb = load_workbook(self.db, read_only=True, data_only=True)
            try:
                ws = wb["DB"] if "DB" in wb.sheetnames else wb.active
                start = core.db_room_start_col(ws)
                entries = []
                for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                    if not any(v not in (None, "") for v in row):
                        continue
                    entries.append(core.DbEntry(index, str(row[0] or "").strip(), str(row[1] or "").strip(),
                        str(row[2] or "").strip(), core.normalize_price(row[3]) if start == 5 else "",
                        [str(v).strip() for v in row[start-1:] if v not in (None, "")]))
                self.cache = (stamp, entries, revision(self.db))
            finally:
                wb.close()
        return self.cache[1], self.cache[2]

    def save_meta(self):
        atomic_json(self.meta_path, self.meta)

    def history(self, key):
        entry = next((e for e in self.entries()[0] if entry_key(e) == key), None)
        if not entry or not self.result.exists():
            return []
        wb = load_workbook(self.result, read_only=True, data_only=True)
        try:
            if entry.sheet_title not in wb.sheetnames:
                return []
            items = core.read_result_history_from_ws(wb[entry.sheet_title])
            return [{"date": v["date_text"], "reserved": v["reserved"], "total": v["total"],
                     "rate": v["reserved"] / v["total"] if v["total"] else None,
                     "price": v["price"], "revenue": (v["reserved"] * exact_price(v["price"]))
                     if exact_price(v["price"]) is not None else None} for v in items]
        finally:
            wb.close()

    def legacy_results(self):
        latest = {}
        for path in sorted(self.backup.glob("booking_run_*.csv"), reverse=True)[:120]:
            if re.fullmatch(r"booking_run_\d{8}_\d{6}_\d{6}\.csv", path.name):
                continue
            try:
                with path.open(encoding="utf-8-sig", newline="") as stream:
                    for row in csv.DictReader(stream):
                        e = core.DbEntry(0, row.get("업체고유번호", ""), "", row.get("중분류", ""), "", [])
                        key = entry_key(e)
                        if key in latest:
                            continue
                        number = lambda col: int(row.get(col) or 0)
                        raw_date = parse_qs(urlparse(row.get("URL", "")).query).get("checkin", [""])[0]
                        checkin = f"{raw_date[:4]}-{raw_date[4:6]}-{raw_date[6:]}" if re.fullmatch(r"\d{8}", raw_date) else ""
                        latest[key] = {"status": row.get("상태", "기록 없음"), "reserved": number("예약객실수"),
                            "matched": number("화면매칭객실수"), "total": number("저장총객실수"),
                            "db_count": number("DB객실수"), "missing": number("누락객실수"),
                            "checked_at": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
                            "legacy": True, "date": checkin, "fingerprint": ""}
            except (ValueError, OSError, UnicodeError):
                continue
        return latest

    def apply(self, payload):
        with core.collection_run_lock(self.folder):
            entries, rev = self.entries()
            if payload.get("revision") != rev:
                raise ValueError("DB가 다른 작업에서 변경되었습니다. 새로고침 후 다시 확인해 주세요.")
            edits = payload.get("edits", [])
            if not edits or len(edits) > 1000:
                raise ValueError("수정할 항목을 선택해 주세요.")
            by_key = {entry_key(e): e for e in entries}
            replacement, added, touched, seen = {}, [], [], set()
            for edit in edits:
                key = edit.get("key")
                if key and (key not in by_key or key in seen):
                    raise ValueError("항목이 변경되었거나 중복 선택되었습니다. 새로고침해 주세요.")
                seen.add(key) if key else None
                row = by_key[key].excel_row if key else max([e.excel_row for e in entries] + [1]) + len(added) + 1
                entry = validate_entry(edit, row)
                if key:
                    if by_key[key].company_id != entry.company_id:
                        raise ValueError("기존 업체번호는 변경할 수 없습니다. 신규 항목으로 추가해 주세요.")
                    replacement[key] = entry
                else:
                    added.append(entry)
                touched.append(entry_key(entry))
            final = [replacement.get(entry_key(e), e) for e in entries] + added
            titles = [e.sheet_title.casefold() for e in final]
            if len(titles) != len(set(titles)):
                raise ValueError("중분류/시트명이 중복됩니다. 분류별로 다른 이름을 입력해 주세요.")
            # Permit repairing old overlaps, but never introduce a new double-counted room.
            def overlaps(items):
                owners = {}
                for e in items:
                    for room in e.rooms:
                        owners.setdefault((e.company_id, core.normalize_name(room)), set()).add(e.excel_row)
                return {(key, tuple(sorted(rows))) for key, rows in owners.items() if len(rows) > 1}
            if overlaps(final) - overlaps(entries):
                raise ValueError("같은 업체의 여러 분류에 동일 객실이 겹칩니다. 객실을 나눠서 저장해 주세요.")
            backup = self.home / ("수정전_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
            backup.mkdir()
            self.save_meta()
            for path in (self.db, self.result, self.meta_path):
                if path.exists():
                    shutil.copy2(path, backup / path.name)
            db_wb = load_workbook(self.db)
            result_wb = load_workbook(self.result) if self.result.exists() else None
            try:
                ws = db_wb["DB"] if "DB" in db_wb.sheetnames else db_wb.active
                if core.db_room_start_col(ws) != 5:
                    ws.insert_cols(4)
                    ws.cell(1, 4, core.PRICE_HEADER)
                for entry in list(replacement.values()) + added:
                    values = [entry.company_id, entry.major, entry.sheet_title, entry.price] + entry.rooms
                    for col in range(1, max(ws.max_column, len(values)) + 1):
                        cell = ws.cell(entry.excel_row, col)
                        cell.value = values[col-1] if col <= len(values) else None
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                for key, entry in replacement.items():
                    old = by_key[key].sheet_title
                    if old != entry.sheet_title and result_wb and old in result_wb.sheetnames:
                        if entry.sheet_title.casefold() in {s.casefold() for s in result_wb.sheetnames}:
                            raise ValueError("변경할 시트명에 과거 기록이 이미 있습니다. 다른 이름을 사용해 주세요.")
                        result_wb[old].title = entry.sheet_title
                    new_key = entry_key(entry)
                    if key != new_key:
                        for field in ("archived", "changed"):
                            self.meta[field] = [new_key if v == key else v for v in self.meta[field]]
                        self.meta["regions"][new_key] = self.meta["regions"].pop(key, "")
                for edit, key in zip(edits, touched):
                    if "region" in edit:
                        self.meta["regions"][key] = str(edit["region"])[:150]
                self.meta["changed"] = sorted(set(self.meta["changed"] + touched))
                core.refresh_db_headers(ws)
                db_temp = self.home / "pending_db.xlsx"
                db_wb.save(db_temp)
                result_temp = self.home / "pending_result.xlsx"
                if result_wb:
                    result_wb.save(result_temp)
                atomic_json(self.home / "transaction.json", {"backup": str(backup)})
                os.replace(db_temp, self.db)
                if result_wb:
                    os.replace(result_temp, self.result)
                self.save_meta()
                (self.home / "transaction.json").unlink()
            except Exception:
                self.recover()
                self.meta = read_json(self.meta_path, self.meta)
                raise
            finally:
                db_wb.close()
                if result_wb:
                    result_wb.close()
                self.cache = None
            return {"keys": touched, "backup": str(backup), "revision": self.entries()[1]}
