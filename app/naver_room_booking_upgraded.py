from __future__ import annotations

import argparse
import csv
import booking_file_lock
import os
import re
import shutil
import sys
import threading
import time
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from room_classifier import filtered_snapshot, partition


APP_TITLE = "네이버 객실 예약현황 자동화 Upgrade"
DB_FILENAME = "객실_DB.xlsx"
RESULT_FILENAME = "예약현황_Result.xlsx"
DEFAULT_STATUS_WORDS = "예매완료,예약마감,예약완료,예약 완료,판매완료,매진"
DEFAULT_FOLDER = os.environ.get(
    "NAVER_BOOKING_FOLDER",
    r"C:\Users\wisec\OneDrive\Desktop\네이버객실예약현황프로그램",
)
BACKUP_FOLDERNAME = "백업"
PRICE_HEADER = "금액"
STATUS_SHEET_NAME = "현황"
RESULT_META_SHEETS = {STATUS_SHEET_NAME}
PRODUCT_TYPE_OPTIONS = ["일반캠핑", "글램핑", "카라반", "펜션", "방갈로", "기타"]
SPECIAL_PRODUCT_TYPES = ["글램핑", "카라반", "펜션", "방갈로"]
LEGACY_ROOM_START_COL = 4
ROOM_START_COL = 5


@dataclass
class DbEntry:
    excel_row: int
    company_id: str
    major: str
    sheet_title: str
    price: str
    rooms: list[str]


@dataclass
class RoomMatch:
    room_name: str
    found: bool
    reserved: bool
    evidence: str


@dataclass
class EntryResult:
    entry: DbEntry
    reserved_count: int
    matched_count: int
    db_room_count: int
    saved_total: int
    missing_count: int
    status: str
    url: str
    reserved_rooms: list[str]
    missing_rooms: list[str]
    detected_rooms: list[str]
    db_updated: bool = False


def log_time() -> str:
    return datetime.now().strftime("%H:%M:%S")


def normalize_name(value: object) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"\s+", "", text)
    text = text.replace("\u200b", "")
    return text.strip()


def normalize_words(words: str | Iterable[str]) -> list[str]:
    if isinstance(words, str):
        raw = [part.strip() for part in words.split(",")]
    else:
        raw = [str(part).strip() for part in words]
    result = []
    for word in raw:
        key = normalize_name(word)
        if key and key not in result:
            result.append(key)
    return result


def contains_room_key(text_key: str, room_key: str) -> bool:
    if not text_key or not room_key:
        return False

    start = text_key.find(room_key)
    while start != -1:
        end = start + len(room_key)
        before = text_key[start - 1] if start > 0 else ""
        after = text_key[end] if end < len(text_key) else ""

        # Avoid A1 matching A10, site1 matching site10, etc.
        if room_key[-1:].isdigit() and after.isdigit():
            start = text_key.find(room_key, start + 1)
            continue

        # Very short room names need ASCII/digit boundaries, but Korean status
        # text right after the room name is allowed: A1예약마감 is a valid card.
        if len(room_key) <= 3:
            if before.isascii() and before.isalnum():
                start = text_key.find(room_key, start + 1)
                continue
            if after.isascii() and after.isalnum():
                start = text_key.find(room_key, start + 1)
                continue

        return True

    return False


def parse_date(value: str | None) -> date:
    if not value:
        return date.today()
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError("날짜는 2026-05-24 또는 20260524 형식으로 입력해 주세요.")


def same_day(cell_value: object, target: date) -> bool:
    if isinstance(cell_value, datetime):
        return cell_value.date() == target
    if isinstance(cell_value, date):
        return cell_value == target
    if isinstance(cell_value, str):
        text = cell_value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y%m%d"):
            try:
                return datetime.strptime(text, fmt).date() == target
            except ValueError:
                pass
    return False


def ensure_folder(folder: str | Path) -> Path:
    path = Path(folder).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"폴더가 존재하지 않습니다: {path}")
    if not path.is_dir():
        raise NotADirectoryError(f"폴더가 아닙니다: {path}")
    return path


def infer_product_type(*parts: object) -> str:
    text = normalize_name(" ".join("" if part is None else str(part) for part in parts)).lower()
    if "글램핑" in text:
        return "글램핑"
    if "카라반" in text or "캠핑카" in text or "트레일러" in text:
        return "카라반"
    if "펜션" in text or "스테이" in text or "풀빌라" in text or "독채" in text:
        return "펜션"
    if "방갈로" in text or "방가로" in text:
        return "방갈로"
    if "오토캠핑" in text or "캠핑" in text or "사이트" in text or "데크" in text or "파쇄석" in text:
        return "일반캠핑"
    return "기타"


def entry_product_type(entry: "DbEntry") -> str:
    title_type = infer_product_type(entry.sheet_title)
    if title_type in SPECIAL_PRODUCT_TYPES:
        return title_type
    major_type = infer_product_type(entry.major)
    if major_type in SPECIAL_PRODUCT_TYPES:
        return major_type
    room_type = infer_product_type(*entry.rooms[:6])
    if room_type in SPECIAL_PRODUCT_TYPES and title_type == "기타" and major_type == "기타":
        return room_type
    return "일반캠핑"


def room_product_type(room_name: object) -> str:
    text = normalize_name(room_name).lower()
    if "글램핑" in text or "glamping" in text:
        return "글램핑"
    if "카라반" in text or "캠핑카" in text or "트레일러" in text or "caravan" in text:
        return "카라반"
    if "펜션" in text or "스테이" in text or "풀빌라" in text or "독채" in text:
        return "펜션"
    if "방갈로" in text or "방가로" in text:
        return "방갈로"
    if "오토캠핑" in text or "캠핑" in text or "사이트" in text or "데크" in text or "파쇄석" in text:
        return "일반캠핑"
    return "기타"


def normalized_room_list(rooms: Iterable[object]) -> list[str]:
    return [normalize_name(room) for room in rooms if normalize_name(room)]


def same_room_list(left: Iterable[object], right: Iterable[object]) -> bool:
    return normalized_room_list(left) == normalized_room_list(right)


def same_room_set(left: Iterable[object], right: Iterable[object]) -> bool:
    left_keys = normalized_room_list(left)
    right_keys = normalized_room_list(right)
    return len(left_keys) == len(right_keys) and set(left_keys) == set(right_keys)


def result_has_full_detected_copy(result: "EntryResult") -> bool:
    return bool(result.detected_rooms) and same_room_set(result.entry.rooms, result.detected_rooms)


def split_rooms_by_product_type(entries: list["DbEntry"], detected_rooms: list[str]) -> tuple[dict[int, list[str]], str]:
    if len(entries) < 2:
        return {}, "같은 업체번호 DB 행이 2개 이상일 때만 자동분리할 수 있습니다."
    if not detected_rooms:
        return {}, "네이버에서 객실명을 찾지 못했습니다."

    type_counts = Counter(entry_product_type(entry) for entry in entries)
    duplicated_types = [product_type for product_type, count in type_counts.items() if count > 1]
    if duplicated_types:
        return {}, f"같은 분류 행이 여러 개 있습니다: {', '.join(duplicated_types)}"

    by_type = {entry_product_type(entry): entry for entry in entries}
    assignments: dict[int, list[str]] = {entry.excel_row: [] for entry in entries}
    fallback_type = "일반캠핑" if "일반캠핑" in by_type else ("기타" if "기타" in by_type else "")
    unassigned: list[str] = []

    for room in detected_rooms:
        detected_type = room_product_type(room)
        target_entry = by_type.get(detected_type)
        if target_entry:
            assignments[target_entry.excel_row].append(room)
        elif detected_type == "기타" and fallback_type:
            assignments[by_type[fallback_type].excel_row].append(room)
        else:
            unassigned.append(room)

    empty_special = [
        entry_product_type(entry)
        for entry in entries
        if entry_product_type(entry) in SPECIAL_PRODUCT_TYPES and not assignments.get(entry.excel_row)
    ]
    if empty_special:
        return {}, (
            "객실명만으로 "
            + ", ".join(empty_special)
            + " 객실을 구분하지 못했습니다. DB관리에서 객실명을 나눠 주세요."
        )
    if unassigned:
        return {}, f"분류를 판단하지 못한 객실 {len(unassigned)}개가 있습니다."

    updates = {
        entry.excel_row: assignments[entry.excel_row]
        for entry in entries
        if assignments[entry.excel_row] and not same_room_list(entry.rooms, assignments[entry.excel_row])
    }
    if not updates:
        return {}, "이미 분류별로 나뉘어 있습니다."

    return updates, f"{sum(len(rooms) for rooms in assignments.values())}개 객실을 자동분리했습니다."


def has_meta_column(ws) -> bool:
    header = normalize_name(ws.cell(row=1, column=LEGACY_ROOM_START_COL).value)
    return header in {normalize_name(PRICE_HEADER), normalize_name("상품유형")}


def db_room_start_col(ws) -> int:
    return ROOM_START_COL if has_meta_column(ws) else LEGACY_ROOM_START_COL


def normalize_price(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text in PRODUCT_TYPE_OPTIONS:
        return ""
    return text


def parse_price_amount(value: object) -> float | None:
    text = str(value or "").replace(",", "").replace(" ", "")
    if re.fullmatch(r"\d+(?:\.\d+)?만(?:원)?", text):
        return float(re.sub(r"만원?$", "", text)) * 10000
    if re.fullmatch(r"\d+(?:\.0+)?원?", text):
        amount = float(text.rstrip("원"))
        return amount if amount >= 1000 or text.endswith("원") else None
    return None


def format_price_text(value: object) -> str:
    amount = parse_price_amount(value)
    if amount is None:
        return "" if value is None else str(value).strip()
    return f"{int(amount):,}"


def format_date_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if value in (None, ""):
        return ""
    return str(value)


def safe_int(value: object, default: int = 0) -> int:
    if value in (None, ""):
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def calc_rate(reserved: int, total: int) -> float | None:
    if total <= 0:
        return None
    return reserved / total


def refresh_db_headers(ws) -> None:
    ws.cell(row=1, column=1).value = "업체 고유인덱스번호"
    ws.cell(row=1, column=2).value = "대분류"
    ws.cell(row=1, column=3).value = "중분류(시트명에 반영)"
    ws.cell(row=1, column=4).value = PRICE_HEADER
    for col in range(ROOM_START_COL, max(ws.max_column, ROOM_START_COL) + 1):
        ws.cell(row=1, column=col).value = f"세분류_{col - ROOM_START_COL + 1}"


def ensure_db_schema(db_path: Path) -> bool:
    if not db_path.exists():
        raise FileNotFoundError(f"{DB_FILENAME} 파일이 없습니다: {db_path}")

    wb = load_workbook(db_path)
    ws = wb["DB"] if "DB" in wb.sheetnames else wb.active
    if has_meta_column(ws):
        if normalize_name(ws.cell(row=1, column=LEGACY_ROOM_START_COL).value) == normalize_name("상품유형"):
            backup_db_file(db_path)
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=4).value in PRODUCT_TYPE_OPTIONS:
                    ws.cell(row=row, column=4).value = None
            refresh_db_headers(ws)
            wb.save(db_path)
        return False

    backup_db_file(db_path)
    ws.insert_cols(LEGACY_ROOM_START_COL)
    refresh_db_headers(ws)
    wb.save(db_path)
    return True


def load_db(db_path: Path, migrate: bool = True) -> list[DbEntry]:
    if not db_path.exists():
        raise FileNotFoundError(f"{DB_FILENAME} 파일이 없습니다: {db_path}")

    if migrate:
        ensure_db_schema(db_path)

    wb = load_workbook(db_path, read_only=True, data_only=True)
    ws = wb["DB"] if "DB" in wb.sheetnames else wb.active
    entries: list[DbEntry] = []
    room_start_col = db_room_start_col(ws)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in row):
            continue

        company_id = "" if row[0] is None else str(row[0]).strip()
        major = "" if row[1] is None else str(row[1]).strip()
        sheet_title = "" if row[2] is None else str(row[2]).strip()
        price = ""
        if room_start_col == ROOM_START_COL and len(row) >= 4:
            price = normalize_price(row[3])
        rooms = [
            str(value).strip()
            for value in row[room_start_col - 1 :]
            if value not in (None, "") and str(value).strip()
        ]

        if not company_id or not sheet_title:
            raise ValueError(f"DB {row_idx}행의 업체 고유번호 또는 중분류가 비어 있습니다.")

        entries.append(
            DbEntry(
                excel_row=row_idx,
                company_id=company_id,
                major=major,
                sheet_title=sheet_title,
                price=price,
                rooms=rooms,
            )
        )

    wb.close()
    if not entries:
        raise ValueError("DB에 수집할 업체가 없습니다.")

    return entries


def ensure_result_workbook(result_path: Path) -> None:
    if result_path.exists():
        return
    wb = Workbook()
    wb.save(result_path)


def assert_result_writable(result_path: Path) -> None:
    ensure_result_workbook(result_path)
    try:
        with open(result_path, "r+b"):
            pass
    except PermissionError as exc:
        raise PermissionError(f"{RESULT_FILENAME} 파일이 열려 있습니다. 엑셀을 닫고 다시 실행해 주세요.") from exc


def backup_result_file(result_path: Path) -> Path | None:
    if not result_path.exists():
        return None
    backup_dir = result_path.parent / BACKUP_FOLDERNAME
    backup_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"예약현황_Result_backup_{stamp}.xlsx"
    shutil.copy2(result_path, backup_path)
    return backup_path


def backup_db_file(db_path: Path) -> Path | None:
    if not db_path.exists():
        return None
    backup_dir = db_path.parent / BACKUP_FOLDERNAME
    backup_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"객실_DB_backup_{stamp}.xlsx"
    shutil.copy2(db_path, backup_path)
    return backup_path


def write_db_rooms(db_path: Path, updates: dict[int, list[str]]) -> int:
    if not updates:
        return 0

    ensure_db_schema(db_path)
    wb = load_workbook(db_path)
    ws = wb["DB"] if "DB" in wb.sheetnames else wb.active
    start_col = db_room_start_col(ws)
    max_room_count = max((len(rooms) for rooms in updates.values()), default=0)
    max_col = max(ws.max_column, start_col + max_room_count + 5)

    for row_idx, rooms in updates.items():
        for col in range(start_col, max_col + 1):
            ws.cell(row=row_idx, column=col).value = None
        for offset, room in enumerate(rooms):
            ws.cell(row=row_idx, column=start_col + offset).value = room

    refresh_db_headers(ws)
    wb.save(db_path)
    return len(updates)


def style_header(ws) -> None:
    headers = ["일자", "예약객실 수", "총 객실수", "가동률", "금액", "추정매출"]
    fill = PatternFill(start_color="5A9BD5", end_color="5A9BD5", fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 16


def get_or_create_result_sheet(wb, title: str):
    if title in wb.sheetnames:
        ws = wb[title]
        if ws.max_row < 1 or ws["A1"].value != "일자":
            style_header(ws)
        return ws

    ws = wb.create_sheet(title=title)
    style_header(ws)
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        blank = wb["Sheet"]
        if blank.max_row == 1 and blank.max_column == 1 and blank["A1"].value is None:
            wb.remove(blank)
    return ws


def write_result_row(wb, result: EntryResult, target_date: date) -> None:
    ws = get_or_create_result_sheet(wb, result.entry.sheet_title)

    found_row = None
    for row in range(2, ws.max_row + 1):
        if same_day(ws[f"A{row}"].value, target_date):
            found_row = row
            break

    write_row = found_row or ws.max_row + 1
    ws[f"A{write_row}"] = target_date
    ws[f"A{write_row}"].number_format = "yyyy-mm-dd"
    ws[f"B{write_row}"] = result.reserved_count
    ws[f"C{write_row}"] = result.saved_total
    rate = calc_rate(result.reserved_count, result.saved_total)
    ws[f"D{write_row}"] = rate if rate is not None else None
    ws[f"D{write_row}"].number_format = "0.0%"
    price = parse_price_amount(result.entry.price)
    ws[f"E{write_row}"] = price if price is not None else result.entry.price
    if price is not None:
        ws[f"E{write_row}"].number_format = "#,##0"
        ws[f"F{write_row}"] = result.reserved_count * price
        ws[f"F{write_row}"].number_format = "#,##0"
    else:
        ws[f"F{write_row}"] = None


def read_result_history_from_ws(ws) -> list[dict[str, object]]:
    history: list[dict[str, object]] = []
    for values in ws.iter_rows(min_row=2, max_col=6, values_only=True):
        day = values[0]
        if day in (None, ""):
            continue
        reserved = safe_int(values[1])
        total = safe_int(values[2])
        rate = values[3]
        if rate in (None, ""):
            rate = calc_rate(reserved, total)
        price = values[4]
        revenue = values[5]
        if revenue in (None, ""):
            amount = parse_price_amount(price)
            revenue = reserved * amount if amount is not None else None
        history.append(
            {
                "date": day,
                "date_text": format_date_text(day),
                "reserved": reserved,
                "total": total,
                "rate": rate,
                "price": price,
                "revenue": revenue,
            }
        )
    history.sort(key=lambda item: item["date_text"])
    return history


def latest_history_item(wb, sheet_title: str) -> dict[str, object] | None:
    if sheet_title not in wb.sheetnames:
        return None
    history = read_result_history_from_ws(wb[sheet_title])
    return history[-1] if history else None


def write_status_sheet(wb, entries: list[DbEntry]) -> None:
    if STATUS_SHEET_NAME in wb.sheetnames:
        ws = wb[STATUS_SHEET_NAME]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(STATUS_SHEET_NAME, 0)

    headers = [
        "상태",
        "DB행",
        "고유인덱스",
        "대분류",
        "중분류",
        "금액",
        "최근일자",
        "예약객실",
        "총객실",
        "가동률",
        "추정매출",
        "비고",
    ]
    fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.fill = fill
        cell.font = font
        cell.alignment = align

    active_titles = {entry.sheet_title for entry in entries}
    row_idx = 2
    for entry in entries:
        latest = latest_history_item(wb, entry.sheet_title)
        status = "정상" if latest else "수집전"
        reserved = safe_int(latest["reserved"]) if latest else 0
        total = safe_int(latest["total"]) if latest else len(entry.rooms)
        rate = calc_rate(reserved, total)
        price = parse_price_amount(entry.price)
        revenue = reserved * price if price is not None else None
        values = [
            status,
            entry.excel_row,
            entry.company_id,
            entry.major,
            entry.sheet_title,
            price if price is not None else entry.price,
            latest["date"] if latest else "",
            reserved if latest else "",
            total,
            rate,
            revenue,
            "",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col).value = value
        ws.cell(row=row_idx, column=7).number_format = "yyyy-mm-dd"
        ws.cell(row=row_idx, column=10).number_format = "0.0%"
        ws.cell(row=row_idx, column=6).number_format = "#,##0"
        ws.cell(row=row_idx, column=11).number_format = "#,##0"
        row_idx += 1

    for sheet_name in wb.sheetnames:
        if sheet_name in RESULT_META_SHEETS or sheet_name in active_titles:
            continue
        ws_old = wb[sheet_name]
        if ws_old.max_row <= 1 and ws_old.max_column <= 1 and ws_old["A1"].value is None:
            continue
        latest = latest_history_item(wb, sheet_name)
        reserved = safe_int(latest["reserved"]) if latest else 0
        total = safe_int(latest["total"]) if latest else 0
        rate = calc_rate(reserved, total)
        values = [
            "DB없음",
            "",
            "",
            "",
            sheet_name,
            "",
            latest["date"] if latest else "",
            reserved if latest else "",
            total if latest else "",
            rate,
            "",
            "객실_DB.xlsx에 없는 옛 결과 시트입니다.",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col).value = value
        ws.cell(row=row_idx, column=7).number_format = "yyyy-mm-dd"
        ws.cell(row=row_idx, column=10).number_format = "0.0%"
        row_idx += 1

    widths = [12, 8, 14, 18, 24, 12, 14, 10, 10, 10, 14, 34]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.freeze_panes = "A2"


def read_status_rows(folder: Path) -> list[dict[str, object]]:
    db_path = folder / DB_FILENAME
    result_path = folder / RESULT_FILENAME
    entries = load_db(db_path)
    entry_by_title = {entry.sheet_title: entry for entry in entries}
    rows: list[dict[str, object]] = []
    if not result_path.exists():
        for entry in entries:
            rows.append({"state": "수집전", "entry": entry, "sheet": entry.sheet_title, "latest": None, "orphan": False})
        return rows

    wb = load_workbook(result_path, data_only=True)
    for entry in entries:
        rows.append(
            {
                "state": "정상" if entry.sheet_title in wb.sheetnames else "수집전",
                "entry": entry,
                "sheet": entry.sheet_title,
                "latest": latest_history_item(wb, entry.sheet_title),
                "orphan": False,
            }
        )
    for sheet_name in wb.sheetnames:
        if sheet_name in RESULT_META_SHEETS or sheet_name in entry_by_title:
            continue
        ws = wb[sheet_name]
        if ws.max_row <= 1 and ws.max_column <= 1 and ws["A1"].value is None:
            continue
        rows.append(
            {
                "state": "DB없음",
                "entry": None,
                "sheet": sheet_name,
                "latest": latest_history_item(wb, sheet_name),
                "orphan": True,
            }
        )
    return rows


def read_result_history(folder: Path, sheet_title: str) -> list[dict[str, object]]:
    result_path = folder / RESULT_FILENAME
    if not result_path.exists():
        return []
    wb = load_workbook(result_path, data_only=True)
    if sheet_title not in wb.sheetnames:
        return []
    return read_result_history_from_ws(wb[sheet_title])


def build_naver_room_url(company_id: str, checkin: date) -> str:
    checkout = checkin + timedelta(days=1)
    return (
        f"https://m.place.naver.com/accommodation/{company_id}/room"
        f"?entry=pll&businessCategory=camping&level=top&guest=1"
        f"&checkin={checkin.strftime('%Y%m%d')}"
        f"&checkout={checkout.strftime('%Y%m%d')}"
    )


def parse_company_id(value: str) -> str:
    text = (value or "").strip()
    if not text:
        raise ValueError("네이버 링크 또는 업체 고유번호를 입력해 주세요.")

    patterns = [
        r"/accommodation/(\d+)",
        r"/place/(\d+)",
        r"/restaurant/(\d+)",
        r"/hospital/(\d+)",
        r"/hairshop/(\d+)",
        r"place/(\d+)",
        r"accommodation/(\d+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)

    digits = re.sub(r"\D", "", text)
    if digits:
        return digits

    raise ValueError("링크에서 업체 고유번호를 찾지 못했습니다.")


def add_or_update_db_entry(
    db_path: Path,
    company_id: str,
    major: str,
    sheet_title: str,
    rooms: list[str],
    price: str = "",
) -> tuple[str, int]:
    if not rooms:
        raise ValueError("네이버 화면에서 객실명을 찾지 못했습니다.")

    ensure_db_schema(db_path)
    wb = load_workbook(db_path)
    ws = wb["DB"] if "DB" in wb.sheetnames else wb.active
    target_row = None
    price = normalize_price(price)

    for row in range(2, ws.max_row + 1):
        row_company_id = "" if ws.cell(row=row, column=1).value is None else str(ws.cell(row=row, column=1).value).strip()
        row_sheet_title = "" if ws.cell(row=row, column=3).value is None else str(ws.cell(row=row, column=3).value).strip()
        if row_company_id == company_id and row_sheet_title == sheet_title:
            target_row = row
            break

    action = "updated"
    if target_row is None:
        target_row = ws.max_row + 1
        action = "added"
        ws.cell(row=target_row, column=1).value = company_id
        ws.cell(row=target_row, column=2).value = major
        ws.cell(row=target_row, column=3).value = sheet_title
    else:
        if major:
            ws.cell(row=target_row, column=2).value = major
        ws.cell(row=target_row, column=3).value = sheet_title
    ws.cell(row=target_row, column=4).value = price

    start_col = db_room_start_col(ws)
    max_col = max(ws.max_column, start_col + len(rooms) + 5)
    for col in range(start_col, max_col + 1):
        ws.cell(row=target_row, column=col).value = None
    for offset, room in enumerate(rooms):
        ws.cell(row=target_row, column=start_col + offset).value = room

    refresh_db_headers(ws)
    wb.save(db_path)
    return action, target_row


def save_db_entry(
    db_path: Path,
    row_idx: int | None,
    company_id: str,
    major: str,
    sheet_title: str,
    price: str,
    rooms: list[str],
) -> int:
    company_id = company_id.strip()
    major = major.strip()
    sheet_title = sheet_title.strip()
    rooms = [room.strip() for room in rooms if room.strip()]
    if not company_id:
        raise ValueError("업체 고유인덱스번호를 입력해 주세요.")
    if not sheet_title:
        raise ValueError("중분류/시트명을 입력해 주세요.")
    if not rooms:
        raise ValueError("세분류 객실명을 1개 이상 입력해 주세요.")

    ensure_db_schema(db_path)
    wb = load_workbook(db_path)
    ws = wb["DB"] if "DB" in wb.sheetnames else wb.active
    target_row = row_idx if row_idx and row_idx >= 2 else ws.max_row + 1
    if target_row > ws.max_row + 1:
        raise ValueError(f"DB {target_row}행을 찾지 못했습니다.")

    ws.cell(row=target_row, column=1).value = company_id
    ws.cell(row=target_row, column=2).value = major
    ws.cell(row=target_row, column=3).value = sheet_title
    ws.cell(row=target_row, column=4).value = normalize_price(price)

    start_col = db_room_start_col(ws)
    max_col = max(ws.max_column, start_col + len(rooms) + 5)
    for col in range(start_col, max_col + 1):
        ws.cell(row=target_row, column=col).value = None
    for offset, room in enumerate(rooms):
        ws.cell(row=target_row, column=start_col + offset).value = room

    refresh_db_headers(ws)
    wb.save(db_path)
    return target_row


def delete_db_entry(db_path: Path, row_idx: int) -> None:
    if row_idx < 2:
        raise ValueError("삭제할 DB 행을 선택해 주세요.")
    ensure_db_schema(db_path)
    wb = load_workbook(db_path)
    ws = wb["DB"] if "DB" in wb.sheetnames else wb.active
    if row_idx > ws.max_row:
        raise ValueError(f"DB {row_idx}행을 찾지 못했습니다.")
    ws.delete_rows(row_idx, 1)
    refresh_db_headers(ws)
    wb.save(db_path)


def split_room_text(value: str) -> list[str]:
    rooms: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[\n\r]+|\s*\|\s*", value or ""):
        room = part.strip()
        key = normalize_name(room)
        if room and key not in seen:
            seen.add(key)
            rooms.append(room)
    return rooms


EXTRACT_SCRIPT = r"""
({ statusKeys }) => {
  const normalize = (value) => (value || "").replace(/\s+/g, "").replace(/\u200b/g, "").trim();
  const isVisible = (el) => {
    if (!el || el.nodeType !== Node.ELEMENT_NODE) return false;
    const style = window.getComputedStyle(el);
    if (!style || style.display === "none" || style.visibility === "hidden" || style.opacity === "0") return false;
    const rect = el.getBoundingClientRect();
    return rect.width >= 8 && rect.height >= 8;
  };
  const textOf = (el) => (el.innerText || el.textContent || "").trim();
  const hasStatus = (key) => statusKeys.some((status) => key.includes(status));
  const pickTitle = (el, fallbackText) => {
    const productImage = el.querySelector("a[href*='/bizes/'][href*='/items/'] img[alt], a[href*='/room/'] img[alt]");
    if (productImage && productImage.alt.trim()) return productImage.alt.trim();
    const named = el.querySelector('span.fjSjb');
    if (named && textOf(named).trim()) return textOf(named).trim();
    const selectors = [
      "span.fjSjb",
      "strong",
      "b",
      "h2",
      "h3",
      "[class*='title']",
      "[class*='name']"
    ];
    for (const selector of selectors) {
      const items = Array.from(el.querySelectorAll(selector));
      for (const item of items) {
        if (!isVisible(item)) continue;
        const value = textOf(item).split("\n").map((line) => line.trim()).filter(Boolean)[0] || "";
        if (value && !/^[0-9,]+원$/.test(normalize(value)) && !statusKeys.includes(normalize(value))) return value.trim();
      }
    }

    for (const line of fallbackText.split("\n").map((line) => line.trim()).filter(Boolean)) {
      if (line && !/^[0-9,]+원$/.test(normalize(line)) && !statusKeys.includes(normalize(line))) return line;
    }
    return "";
  };
  const candidates = [];
  const roomCards = [];
  const seen = new Set();
  const seenRooms = new Set();
  const roomEvidence = (el, text) => {
    if (el.closest("[class*='review'], [id*='review']")) return false;
    const roomLink = el.querySelector("a[href*='/room/'], a[href*='/bizes/'][href*='/items/']");
    const capacity = /(?:기준|최대)\s*\d+\s*인|체크인|체크아웃|\d+\s*(?:㎡|m²)/.test(text);
    const roomTitle = el.querySelector("span.fjSjb");
    return !!roomLink || capacity || (!!roomTitle && hasStatus(normalize(text)));
  };

  const elements = Array.from(document.querySelectorAll("li, article, [role='listitem']"));

  for (const el of elements) {
    if (!isVisible(el)) continue;
    const text = textOf(el);
    if (!text || text.length > 2500) continue;

    const rect = el.getBoundingClientRect();
    const key = `${Math.round(rect.x)}:${Math.round(rect.y)}:${normalize(text).slice(0, 140)}`;
    if (seen.has(key)) continue;
    seen.add(key);

    if (el.matches("div.place_section_content > div > ul > li, li, article, [role='listitem']") && roomEvidence(el, text)) {
      const nested = Array.from(el.querySelectorAll("li,article,[role='listitem']"));
      if (nested.some(child => child.querySelector("strong,b,h2,h3,span.fjSjb") && roomEvidence(child,textOf(child)))) continue;
      const title = pickTitle(el, text);
      const titleKey = normalize(title);
      const links = Array.from(el.querySelectorAll("a[href*='/bizes/'][href*='/items/'],a[href*='/room/']"));
      const productIds = [...new Set(links.map(a => {const u=new URL(a.href);return u.pathname;}).filter(Boolean))];
      if (productIds.length > 1) continue;
      const productKey = productIds[0] ? 'product:' + productIds[0] : 'name:' + titleKey.toLowerCase();
      const statuses = Array.from(el.querySelectorAll('*')).filter(n=>!n.children.length).map(textOf).filter(t=>statusKeys.includes(normalize(t)));
      const firstLine = text.split('\n').map(s=>s.trim()).find(Boolean) || '';
      if (statusKeys.includes(normalize(firstLine))) statuses.push(firstLine);
      if (title && !seenRooms.has(productKey)) {
        seenRooms.add(productKey);
        roomCards.push({
          schemaVersion: 3,
          itemKey: productKey,
          itemUrl: links[0] ? links[0].href : '',
          unitEvidence: true,
          overnightEvidence: /최소\s*\d+\s*박|\d+\s*박|침실|침대|캠핑사이트|파쇄석|일반노지|글램핑|카라반/.test(text),
          statusLabels: [...new Set(statuses)],
          title,
          titleKey,
          text,
          textKey: normalize(text),
          reserved: statuses.length > 0,
          x: Math.round(rect.x),
          y: Math.round(rect.y),
        });
      }
    }

    if (roomCards.length >= 1200) break;
  }

  const bodyText = (document.body && (document.body.innerText || document.body.textContent)) || "";
  const bodyKey = normalize(bodyText);
  return {
    title: document.title,
    url: location.href,
    bodyKey,
    bodyText: bodyText.slice(0, 8000),
    candidates,
    roomCards,
    roomEvidenceFiltered: true,
    strictRoomCards: true,
    statusOccurrences: statusKeys.reduce((sum, status) => {
      return sum + (bodyKey.split(status).length - 1);
    }, 0),
  };
}
"""


def import_playwright():
    try:
        from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RuntimeError(
            "Playwright가 설치되어 있지 않습니다. install_upgrade.bat를 먼저 실행해 주세요."
        ) from exc
    return sync_playwright, PlaywrightTimeoutError


class BrowserCollector:
    def __init__(
        self,
        checkin: date,
        status_words: list[str],
        show_browser: bool,
        wait_seconds: int,
        progress: Callable[[str], None],
    ) -> None:
        self.checkin = checkin
        self.status_words = status_words
        self.show_browser = show_browser
        self.wait_seconds = wait_seconds
        self.progress = progress
        self._playwright = None
        self._browser = None
        self._context = None
        self._page = None
        self._timeout_error = None
        self._snapshots: dict[str, dict] = {}

    def __enter__(self) -> "BrowserCollector":
        sync_playwright, timeout_error = import_playwright()
        self._timeout_error = timeout_error
        self._playwright = sync_playwright().start()
        self._browser = self._playwright.chromium.launch(
            headless=not self.show_browser,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--disable-extensions",
                "--disable-geolocation",
                "--no-sandbox",
                "--window-size=390,800",
            ],
        )
        self._context = self._browser.new_context(
            viewport={"width": 390, "height": 800},
            locale="ko-KR",
            user_agent=(
                "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) "
                "AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 "
                "Mobile/15E148 Safari/604.1"
            ),
        )
        self._page = self._context.new_page()
        self._page.on("dialog", lambda dialog: dialog.accept())
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        try:
            if self._context:
                self._context.close()
            if self._browser:
                self._browser.close()
        finally:
            if self._playwright:
                self._playwright.stop()

    def get_snapshot(self, company_id: str) -> tuple[str, dict]:
        if company_id in self._snapshots:
            return build_naver_room_url(company_id, self.checkin), self._snapshots[company_id]

        if not self._page:
            raise RuntimeError("브라우저가 준비되지 않았습니다.")

        url = build_naver_room_url(company_id, self.checkin)
        self.progress(f"[{log_time()}] 접속: {company_id}")
        timeout_ms = max(self.wait_seconds, 5) * 1000
        self._page.goto(url, wait_until="domcontentloaded", timeout=timeout_ms)
        response_text = self._page.inner_text("body")
        if any(word in response_text for word in ("서비스 이용이 제한", "접근이 제한", "자동입력 방지", "비정상적인 접근")):
            raise RuntimeError("네이버 접근 제한: DB를 변경하지 않고 중단합니다.")
        try:
            self._page.wait_for_load_state("networkidle", timeout=timeout_ms)
        except self._timeout_error:
            pass
        self._page.wait_for_timeout(2000)
        self._scroll_to_bottom()
        snapshot = self._page.evaluate(EXTRACT_SCRIPT, {"statusKeys": self.status_words})
        self._snapshots[company_id] = snapshot
        return url, snapshot

    def _scroll_to_bottom(self) -> None:
        if not self._page:
            return
        last_height = 0
        stable_count = 0
        for _ in range(12):
            try:
                new_height = self._page.evaluate("document.body ? document.body.scrollHeight : 0")
                self._page.evaluate("window.scrollTo(0, document.body ? document.body.scrollHeight : 0)")
                self._page.wait_for_timeout(1200)
            except Exception:
                break
            if new_height == last_height:
                stable_count += 1
                if stable_count >= 2:
                    break
            else:
                stable_count = 0
                last_height = new_height


def room_key_positions(text_key: str, room_key: str) -> list[int]:
    positions: list[int] = []
    start = text_key.find(room_key)
    while start != -1:
        end = start + len(room_key)
        before = text_key[start - 1] if start > 0 else ""
        after = text_key[end] if end < len(text_key) else ""
        valid = True
        if room_key[-1:].isdigit() and after.isdigit():
            valid = False
        if len(room_key) <= 3:
            if before.isascii() and before.isalnum():
                valid = False
            if after.isascii() and after.isalnum():
                valid = False
        if valid:
            positions.append(start)
        start = text_key.find(room_key, start + 1)
    return positions


def line_contains_any_room(line_key: str, room_keys: list[str], current_key: str) -> bool:
    return any(key != current_key and contains_room_key(line_key, key) for key in room_keys)


def source_contains_room_key(source_text: object, room_key: str) -> bool:
    text = "" if source_text is None else str(source_text)
    for line in re.split(r"[\r\n]+", text):
        if contains_room_key(normalize_name(line), room_key):
            return True
    return contains_room_key(normalize_name(text), room_key)


def status_near_room_key(
    room_key: str,
    source_text: object,
    status_words: list[str],
    room_keys: list[str],
) -> bool:
    text = "" if source_text is None else str(source_text)
    if not text or not room_key:
        return False

    # Prefer visible line order: a status belongs only to the current room
    # until another known room name starts the next block.
    lines = [line.strip() for line in re.split(r"[\r\n]+", text) if line.strip()]
    for index, line in enumerate(lines):
        line_key = normalize_name(line)
        if not contains_room_key(line_key, room_key):
            continue

        segment = [line]
        for next_line in lines[index + 1 : index + 7]:
            next_key = normalize_name(next_line)
            if line_contains_any_room(next_key, room_keys, room_key):
                break
            segment.append(next_line)
            if len(normalize_name(" ".join(segment))) >= 260:
                break

        segment_key = normalize_name(" ".join(segment))
        if any(status in segment_key for status in status_words):
            return True

    text_key = normalize_name(text)
    positions = room_key_positions(text_key, room_key)
    if not positions:
        return False

    all_other_positions: list[int] = []
    for key in room_keys:
        if key == room_key:
            continue
        all_other_positions.extend(room_key_positions(text_key, key))

    for start in positions:
        following = [pos for pos in all_other_positions if pos > start]
        next_room = min(following) if following else len(text_key)
        end = min(next_room, start + 260)
        begin = max(0, start - 40)
        segment_key = text_key[begin:end]
        if any(status in segment_key for status in status_words):
            return True

    return False


def choose_best_candidate(
    room_key: str,
    candidates: list[dict],
    status_words: list[str],
    room_keys: list[str],
) -> tuple[dict | None, bool]:
    scored = []
    for candidate in candidates:
        text_key = candidate.get("textKey") or ""
        context_key = candidate.get("contextKey") or ""
        if not (
            source_contains_room_key(candidate.get("text"), room_key)
            or source_contains_room_key(candidate.get("context"), room_key)
            or contains_room_key(text_key, room_key)
            or contains_room_key(context_key, room_key)
        ):
            continue

        reserved = status_near_room_key(
            room_key,
            "\n".join([str(candidate.get("text") or ""), str(candidate.get("context") or "")]),
            status_words,
            room_keys,
        )
        text_len = len(candidate.get("text") or "")
        context_len = len(candidate.get("context") or "")
        score = context_len + (text_len * 0.3)
        if candidate.get("preferred"):
            score -= 700
        if reserved:
            score -= 400
        scored.append((score, candidate, reserved))

    if not scored:
        return None, False

    scored.sort(key=lambda item: item[0])
    _, best, reserved = scored[0]
    return best, reserved


def choose_best_room_card(
    room_key: str,
    room_cards: list[dict],
    status_words: list[str],
    room_keys: list[str],
) -> tuple[dict | None, bool]:
    scored = []
    for card in room_cards:
        title_key = card.get("titleKey") or normalize_name(card.get("title"))
        text_key = card.get("textKey") or normalize_name(card.get("text"))
        if not title_key:
            continue
        if not contains_room_key(title_key, room_key):
            continue

        title_gap = abs(len(title_key) - len(room_key))
        text_len = len(card.get("text") or "")
        y = int(card.get("y") or 0)
        score = title_gap * 20 + min(text_len, 2000) * 0.05 + y * 0.001
        if title_key == room_key:
            score -= 200
        reserved = any(normalize_name(label) in status_words for label in card.get("statusLabels", [])) if card.get("schemaVersion") == 3 else status_near_room_key(
            room_key,
            "\n".join([str(card.get("title") or ""), str(card.get("text") or "")]),
            status_words,
            room_keys,
        )
        scored.append((score, card, reserved))

    if not scored:
        return None, False

    scored.sort(key=lambda item: item[0])
    exact = [item for item in scored if normalize_name(item[1].get("title")) == room_key]
    if len(exact) == 1:
        _, best, reserved = exact[0]
        return best, reserved
    if len(scored) > 1:
        return None, False
    _, best, reserved = scored[0]
    return best, reserved


def clean_detected_room_name(value: object, status_words: list[str]) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return ""

    key = normalize_name(text)
    noise = {
        "예약",
        "예약하기",
        "예약가능",
        "예약 가능",
        "상세보기",
        "상세 보기",
        "더보기",
        "펼쳐보기",
        "접기",
        "객실",
        "사이트",
        "날짜선택",
        "가격",
        "요금",
    }
    if text in noise or key in {normalize_name(item) for item in noise}:
        return ""
    noise_keywords = (
        "쿠폰",
        "할인",
        "혜택",
        "이벤트",
        "프로모션",
        "네이버페이",
        "포인트",
        "적립",
        "리뷰",
        "사진",
        "알림받기",
        "공유",
        "찜하기",
        "결제",
        "무료",
        "공지",
        "안내",
    )
    if any(keyword in key for keyword in noise_keywords):
        return ""
    if re.search(r"\*{2,}|입장권|이용권|오전권|오후권|바베큐장|바비큐장", key):
        return ""
    if re.match(r"^(기준|최대)\d+인|^(체크인|체크아웃)", key):
        return ""
    if any(status in key for status in status_words):
        return ""
    if re.fullmatch(r"[0-9,]+(원|%|개)?", key):
        return ""
    if len(text) > 180:
        return ""
    return text


def extract_detected_rooms(snapshot: dict, status_words: list[str]) -> list[str]:
    rooms: list[str] = []
    seen: set[str] = set()

    for card in partition(snapshot, snapshot.get("roomRules"))["room"]:
        room = str(card.get("title") or "").strip()
        key = normalize_name(room)
        if room and key not in seen:
            seen.add(key)
            rooms.append(room)

    # Fallback: old Naver pages often expose the room title as span.fjSjb.
    if not rooms and not snapshot.get("roomEvidenceFiltered"):
        for candidate in snapshot.get("candidates") or []:
            if candidate.get("tag") != "span":
                continue
            room = clean_detected_room_name(candidate.get("text"), status_words)
            key = normalize_name(room)
            if room and key not in seen:
                seen.add(key)
                rooms.append(room)

    return rooms


def analyze_entry_snapshot(entry: DbEntry, snapshot: dict, status_words: list[str], total_mode: str, url: str) -> EntryResult:
    snapshot, _ = filtered_snapshot(snapshot, snapshot.get("roomRules"))
    candidates = snapshot.get("candidates") or []
    room_cards = snapshot.get("roomCards") or []
    detected_rooms = extract_detected_rooms(snapshot, status_words)
    matches: list[RoomMatch] = []
    room_keys = [normalize_name(room) for room in entry.rooms if normalize_name(room)]

    for room in entry.rooms:
        room_key = normalize_name(room)
        card, reserved = choose_best_room_card(room_key, room_cards, status_words, room_keys)
        if card:
            evidence = " ".join((card.get("text") or card.get("title") or "").split())
            matches.append(RoomMatch(room, True, reserved, evidence[:300]))
            continue

        candidate, reserved = (None, False) if snapshot.get("strictRoomCards") else choose_best_candidate(room_key, candidates, status_words, room_keys)
        if candidate:
            evidence = " ".join((candidate.get("context") or candidate.get("text") or "").split())
            matches.append(RoomMatch(room, True, reserved, evidence[:300]))
        else:
            matches.append(RoomMatch(room, False, False, ""))

    matched = [match for match in matches if match.found]
    reserved = [match for match in matched if match.reserved]
    missing = [match.room_name for match in matches if not match.found]

    matched_count = len(matched)
    db_count = len(entry.rooms)
    if total_mode == "db":
        saved_total = db_count
    elif total_mode == "max":
        saved_total = max(db_count, matched_count)
    else:
        saved_total = matched_count

    if db_count == 0:
        status = "DB 객실 없음"
    elif matched_count == 0:
        status_occurrences = snapshot.get("statusOccurrences") or 0
        status = f"객실명 매칭 실패, DB후보 {len(detected_rooms)}개, 상태문구 {status_occurrences}회 발견"
    elif missing:
        status = f"일부 매칭 실패 {len(missing)}개, DB후보 {len(detected_rooms)}개"
    else:
        status = "정상"

    return EntryResult(
        entry=entry,
        reserved_count=len(reserved),
        matched_count=matched_count,
        db_room_count=db_count,
        saved_total=saved_total,
        missing_count=len(missing),
        status=status,
        url=url,
        reserved_rooms=[match.room_name for match in reserved],
        missing_rooms=missing,
        detected_rooms=detected_rooms,
    )


def write_run_log(log_path: Path, results: list[EntryResult]) -> None:
    log_path.parent.mkdir(exist_ok=True)
    with open(log_path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(
            [
                "엑셀행",
                "업체고유번호",
                "대분류",
                "중분류",
                "금액",
                "예약객실수",
                "화면매칭객실수",
                "DB객실수",
                "저장총객실수",
                "누락객실수",
                "상태",
                "URL",
                "예약객실명",
                "누락객실명",
                "네이버감지객실명",
                "DB업데이트",
            ]
        )
        for result in results:
            writer.writerow(
                [
                    result.entry.excel_row,
                    result.entry.company_id,
                    result.entry.major,
                    result.entry.sheet_title,
                    result.entry.price,
                    result.reserved_count,
                    result.matched_count,
                    result.db_room_count,
                    result.saved_total,
                    result.missing_count,
                    result.status,
                    result.url,
                    " | ".join(result.reserved_rooms),
                    " | ".join(result.missing_rooms[:40]),
                    " | ".join(result.detected_rooms[:80]),
                    "Y" if result.db_updated else "",
                ]
            )


def db_rooms_changed(result: EntryResult) -> bool:
    return not same_room_list(result.detected_rooms, result.entry.rooms)


def result_needs_db_repair(result: EntryResult) -> bool:
    return (
        result.missing_count > 0
        or result.matched_count == 0
        or result.status.startswith("오류:")
        or result.status.startswith("분류필요")
    )


def write_db_update_candidates(log_path: Path, results: list[EntryResult]) -> int:
    company_counts = Counter(result.entry.company_id for result in results)
    candidates = [
        result
        for result in results
        if result.detected_rooms
        and (
            result_needs_db_repair(result)
            or (company_counts[result.entry.company_id] == 1 and db_rooms_changed(result))
        )
    ]
    if not candidates:
        return 0

    log_path.parent.mkdir(exist_ok=True)
    with open(log_path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(
            [
                "엑셀행",
                "업체고유번호",
                "대분류",
                "중분류",
                "금액",
                "현재DB객실수",
                "네이버감지객실수",
                "자동업데이트여부",
                "관리메모",
                "현재DB객실명",
                "네이버감지객실명",
            ]
        )
        for result in candidates:
            writer.writerow(
                [
                    result.entry.excel_row,
                    result.entry.company_id,
                    result.entry.major,
                    result.entry.sheet_title,
                    result.entry.price,
                    result.db_room_count,
                    len(result.detected_rooms),
                    "Y" if result.db_updated else "",
                    "같은 업체번호가 여러 행이면 자동으로 섞지 않습니다. DB관리에서 분류해 주세요."
                    if company_counts[result.entry.company_id] > 1
                    else "",
                    " | ".join(result.entry.rooms),
                    " | ".join(result.detected_rooms),
                ]
            )
    return len(candidates)


def update_single_campground_db(
    folder: Path,
    target: str,
    major: str,
    sheet_title: str,
    price: str,
    checkin: date,
    status_words: list[str],
    show_browser: bool,
    wait_seconds: int,
    progress: Callable[[str], None],
) -> tuple[str, int, int]:
    company_id = parse_company_id(target)
    major = major.strip()
    sheet_title = sheet_title.strip()
    if not major:
        major = sheet_title or company_id
    if not sheet_title:
        sheet_title = major
    price = normalize_price(price)

    db_path = folder / DB_FILENAME
    if not db_path.exists():
        raise FileNotFoundError(f"{DB_FILENAME} 파일이 없습니다: {db_path}")

    with BrowserCollector(checkin, status_words, show_browser, wait_seconds, progress) as collector:
        url, snapshot = collector.get_snapshot(company_id)
        rooms = extract_detected_rooms(snapshot, status_words)

    if not rooms:
        raise ValueError("네이버 화면에서 객실명을 찾지 못했습니다. 날짜를 바꾸거나 브라우저 화면을 확인해 주세요.")

    db_backup_path = backup_db_file(db_path)
    if db_backup_path:
        progress(f"[{log_time()}] DB 파일 백업 완료: {db_backup_path.name}")

    action, row = add_or_update_db_entry(db_path, company_id, major, sheet_title, rooms, price)
    action_text = "추가" if action == "added" else "업데이트"
    progress(f"[{log_time()}] DB {action_text} 완료: {sheet_title} / 금액 {price or '-'} ({len(rooms)}개 객실, {row}행)")

    candidate_path = folder / BACKUP_FOLDERNAME / f"single_db_update_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    candidate_path.parent.mkdir(exist_ok=True)
    with open(candidate_path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        room_headers = [f"세분류_{index}" for index in range(1, len(rooms) + 1)]
        writer.writerow(["업체 고유인덱스번호", "대분류", "중분류(시트명에 반영)", "금액", *room_headers])
        writer.writerow([company_id, major, sheet_title, price, *rooms])

    detail_path = folder / BACKUP_FOLDERNAME / f"single_db_update_detail_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    with open(detail_path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(["업체고유번호", "대분류", "중분류", "금액", "객실수", "URL", "객실명"])
        writer.writerow([company_id, major, sheet_title, price, len(rooms), url, " | ".join(rooms)])
    progress(f"[{log_time()}] 추가/업데이트 기록 저장: {candidate_path}")
    return action, row, len(rooms)


@contextmanager
def collection_run_lock(folder: Path):
    lock_path = folder / BACKUP_FOLDERNAME / "booking_collection.lock"
    lock_path.parent.mkdir(exist_ok=True)
    lock_file = open(lock_path, "a+b")
    locked = False
    try:
        lock_file.seek(0, os.SEEK_END)
        if lock_file.tell() == 0:
            lock_file.write(b"\0")
            lock_file.flush()
        lock_file.seek(0)
        try:
            booking_file_lock.acquire(lock_file)
            locked = True
        except OSError as exc:
            raise RuntimeError("다른 예약현황 수집이 이미 실행 중입니다. 기존 수집이 끝난 뒤 다시 실행해 주세요.") from exc

        lock_file.seek(1)
        lock_file.truncate()
        lock_file.write(f"pid={os.getpid()} started={datetime.now().isoformat()}".encode("ascii"))
        lock_file.flush()
        yield
    finally:
        if locked:
            lock_file.seek(0)
            booking_file_lock.release(lock_file)
        lock_file.close()


def run_collection(
    folder: Path,
    checkin: date,
    status_words: list[str],
    total_mode: str,
    db_update_mode: str,
    show_browser: bool,
    wait_seconds: int,
    progress: Callable[[str], None],
    stop_event: threading.Event | None = None,
    result_callback: Callable[[EntryResult], None] | None = None,
    target_excel_rows: set[int] | None = None,
) -> list[EntryResult]:
    with collection_run_lock(folder):
        return _run_collection_impl(
            folder=folder,
            checkin=checkin,
            status_words=status_words,
            total_mode=total_mode,
            db_update_mode=db_update_mode,
            show_browser=show_browser,
            wait_seconds=wait_seconds,
            progress=progress,
            stop_event=stop_event,
            result_callback=result_callback,
            target_excel_rows=target_excel_rows,
        )


def _run_collection_impl(
    folder: Path,
    checkin: date,
    status_words: list[str],
    total_mode: str,
    db_update_mode: str,
    show_browser: bool,
    wait_seconds: int,
    progress: Callable[[str], None],
    stop_event: threading.Event | None = None,
    result_callback: Callable[[EntryResult], None] | None = None,
    target_excel_rows: set[int] | None = None,
) -> list[EntryResult]:
    db_path = folder / DB_FILENAME
    result_path = folder / RESULT_FILENAME
    all_entries = load_db(db_path)
    if target_excel_rows is not None:
        entries = [entry for entry in all_entries if entry.excel_row in target_excel_rows]
    else:
        entries = all_entries
    company_counts = Counter(entry.company_id for entry in all_entries)
    entries_by_company: dict[str, list[DbEntry]] = {}
    for entry in all_entries:
        entries_by_company.setdefault(entry.company_id, []).append(entry)
    assert_result_writable(result_path)
    backup_path = backup_result_file(result_path)

    if backup_path:
        progress(f"[{log_time()}] 결과 파일 백업 완료: {backup_path.name}")
    if target_excel_rows:
        progress(f"[{log_time()}] DB {len(all_entries)}개 행 중 문제 항목 {len(entries)}개만 재수집")
    else:
        progress(f"[{log_time()}] DB {len(entries)}개 행 로드 완료")

    wb = load_workbook(result_path)
    results: list[EntryResult] = []

    with BrowserCollector(checkin, status_words, show_browser, wait_seconds, progress) as collector:
        for index, entry in enumerate(entries, start=1):
            if stop_event and stop_event.is_set():
                progress(f"[{log_time()}] 사용자가 중지했습니다.")
                break
            if not entry.rooms:
                progress(f"[{log_time()}] {entry.excel_row}행 {entry.sheet_title}: 객실명 없음, 건너뜀")
                continue

            try:
                url, snapshot = collector.get_snapshot(entry.company_id)
                result = analyze_entry_snapshot(entry, snapshot, status_words, total_mode, url)
                if company_counts[entry.company_id] > 1 and result.detected_rooms and result_has_full_detected_copy(result):
                    result.status = (
                        f"분류필요: {entry_product_type(entry)} 행이 네이버 전체객실 {len(result.detected_rooms)}개를 모두 들고 있습니다. "
                        "같은 업체 자동분리 또는 DB관리에서 나눠 주세요"
                    )
                elif company_counts[entry.company_id] > 1 and result_needs_db_repair(result) and result.detected_rooms:
                    result.status = (
                        f"분류필요: 같은 업체번호가 {company_counts[entry.company_id]}개 행에 있음, "
                        f"DB관리에서 {entry_product_type(entry)} 객실만 나눠 주세요"
                    )
                write_result_row(wb, result, checkin)
                wb.save(result_path)
                results.append(result)
                if result_callback:
                    result_callback(result)
                progress(
                    f"[{log_time()}] {index}/{len(entries)} {entry.sheet_title}: "
                    f"{result.reserved_count}/{result.saved_total} (금액 {entry.price or '-'}) "
                    f"(매칭 {result.matched_count}, DB {result.db_room_count}) - {result.status}"
                )
            except Exception as exc:
                result = EntryResult(
                    entry=entry,
                    reserved_count=0,
                    matched_count=0,
                    db_room_count=len(entry.rooms),
                    saved_total=0,
                    missing_count=len(entry.rooms),
                    status=f"오류: {exc}",
                    url=build_naver_room_url(entry.company_id, checkin),
                    reserved_rooms=[],
                    missing_rooms=entry.rooms,
                    detected_rooms=[],
                )
                results.append(result)
                if result_callback:
                    result_callback(result)
                progress(f"[{log_time()}] {entry.sheet_title}: 오류 - {exc}")

    db_updates: dict[int, list[str]] = {}
    if db_update_mode in {"auto-safe", "auto-all"}:
        for company_id, group_entries in entries_by_company.items():
            if len(group_entries) < 2:
                continue
            group_results = [result for result in results if result.entry.company_id == company_id and result.detected_rooms]
            if not group_results:
                continue
            if not any(result.status.startswith("분류필요") for result in group_results):
                continue
            detected_rooms = max((result.detected_rooms for result in group_results), key=len)
            split_updates, split_note = split_rooms_by_product_type(group_entries, detected_rooms)
            if split_updates:
                db_updates.update(split_updates)
                for result in group_results:
                    if result.entry.excel_row in split_updates:
                        result.db_updated = True
                        result.status = f"{result.status} - 자동분리 적용"
                progress(f"[{log_time()}] 같은 업체 자동분리: {company_id} - {split_note}")
            else:
                progress(f"[{log_time()}] 자동분리 보류: {company_id} - {split_note}")

        for result in results:
            unique_company = company_counts[result.entry.company_id] == 1
            changed = db_rooms_changed(result)
            if result.entry.excel_row in db_updates:
                continue
            if unique_company and result.detected_rooms and changed:
                db_updates[result.entry.excel_row] = result.detected_rooms
                result.db_updated = True
            elif not unique_company and result.detected_rooms and changed and result_needs_db_repair(result):
                progress(
                    f"[{log_time()}] 자동보류: {result.entry.sheet_title}은 같은 업체번호가 여러 행이라 "
                    "일반/글램핑이 섞이지 않도록 DB관리에서 분류해 주세요."
                )

        if db_updates:
            db_backup_path = backup_db_file(db_path)
            if db_backup_path:
                progress(f"[{log_time()}] DB 파일 백업 완료: {db_backup_path.name}")
            updated_count = write_db_rooms(db_path, db_updates)
            mode_text = "전체 자동" if db_update_mode == "auto-all" else "안전 자동"
            progress(f"[{log_time()}] DB {mode_text} 업데이트 완료: {updated_count}개 행")
        else:
            progress(f"[{log_time()}] DB 자동 업데이트 대상 없음")

    log_path = folder / BACKUP_FOLDERNAME / f"booking_run_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    write_run_log(log_path, results)
    progress(f"[{log_time()}] 상세 로그 저장: {log_path}")
    if db_update_mode in {"candidates", "auto-safe", "auto-all"}:
        candidate_path = folder / BACKUP_FOLDERNAME / f"db_update_candidates_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        candidate_count = write_db_update_candidates(candidate_path, results)
        if candidate_count:
            progress(f"[{log_time()}] DB 업데이트 후보 저장: {candidate_path}")
    write_status_sheet(wb, load_db(db_path))
    wb.save(result_path)
    progress(f"[{log_time()}] 예약현황 현황 시트 업데이트 완료")
    return results


def run_collection_until_normal(
    folder: Path,
    checkin: date,
    status_words: list[str],
    total_mode: str,
    show_browser: bool,
    wait_seconds: int,
    progress: Callable[[str], None],
    stop_event: threading.Event | None = None,
    max_rounds: int = 5,
    result_callback: Callable[[int, EntryResult], None] | None = None,
    initial_target_rows: set[int] | None = None,
) -> list[EntryResult]:
    max_rounds = max(1, int(max_rounds))
    merged_results: dict[int, EntryResult] = {}
    target_rows: set[int] | None = set(initial_target_rows) if initial_target_rows else None

    for round_index in range(1, max_rounds + 1):
        if stop_event and stop_event.is_set():
            progress(f"[{log_time()}] 자동 일괄 보정 중지")
            break

        if target_rows:
            progress(f"[{log_time()}] 자동 일괄 보정 {round_index}/{max_rounds}회차 시작 - 문제 항목 {len(target_rows)}개만 재수집")
        else:
            progress(f"[{log_time()}] 자동 일괄 보정 {round_index}/{max_rounds}회차 시작 - 전체 수집")

        round_results = run_collection(
            folder=folder,
            checkin=checkin,
            status_words=status_words,
            total_mode=total_mode,
            db_update_mode="auto-all",
            show_browser=show_browser,
            wait_seconds=wait_seconds,
            progress=progress,
            stop_event=stop_event,
            result_callback=(lambda result, round_no=round_index: result_callback(round_no, result))
            if result_callback
            else None,
            target_excel_rows=target_rows,
        )

        for result in round_results:
            merged_results[result.entry.excel_row] = result

        issues = [result for result in merged_results.values() if result_needs_db_repair(result)]
        round_issues = [result for result in round_results if result_needs_db_repair(result)]
        updated = [result for result in round_results if result.db_updated]

        if not issues:
            progress(f"[{log_time()}] 모든 항목 정상 확인")
            break

        progress(
            f"[{log_time()}] 정상 아님: 전체 {len(issues)}개, "
            f"이번 회차 문제 {len(round_issues)}개, DB 자동수정 {len(updated)}개"
        )

        if not updated:
            progress(f"[{log_time()}] 더 이상 자동수정할 DB후보가 없어 반복을 멈춥니다.")
            break

        target_rows = {result.entry.excel_row for result in issues}
        if round_index < max_rounds:
            progress(f"[{log_time()}] 수정된 DB로 문제 항목 {len(target_rows)}개만 다시 수집합니다.")
        else:
            progress(f"[{log_time()}] 최대 반복 횟수에 도달했습니다.")

    return [merged_results[row] for row in sorted(merged_results)]


def cli_main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=APP_TITLE)
    parser.add_argument("--folder", default=DEFAULT_FOLDER, help="객실_DB.xlsx와 예약현황_Result.xlsx가 있는 폴더")
    parser.add_argument("--date", default=None, help="체크인 날짜. 예: 2026-05-24")
    parser.add_argument("--status-words", default=DEFAULT_STATUS_WORDS, help="예약 완료로 볼 문구. 쉼표로 구분")
    parser.add_argument(
        "--total-mode",
        choices=["matched", "db", "max"],
        default="matched",
        help="결과 파일의 총 객실수 기준: matched=화면매칭, db=DB, max=둘 중 큰 값",
    )
    parser.add_argument(
        "--db-update-mode",
        choices=["off", "candidates", "auto-safe", "auto-all"],
        default="candidates",
        help="DB 보정: off=안함, candidates=후보 저장, auto-safe/auto-all=단일 업체번호 자동 반영",
    )
    parser.add_argument("--update-db-target", help="DB에 추가/업데이트할 네이버 링크 또는 업체 고유번호")
    parser.add_argument("--update-db-major", default="", help="신규/특정 캠핑장 대분류")
    parser.add_argument("--update-db-sheet-title", default="", help="신규/특정 캠핑장 중분류/시트명")
    parser.add_argument("--update-db-price", default="", help="금액. 예: 50000 또는 50,000")
    parser.add_argument("--headless", action="store_true", help="브라우저를 화면에 보이지 않게 실행")
    parser.add_argument("--wait-seconds", type=int, default=25)
    parser.add_argument("--auto-repair-rounds", type=int, default=0, help="전체 수집 후 DB 일괄 보정/재수집 반복 횟수")
    parser.add_argument("--gui", action="store_true", help="GUI로 실행")
    args = parser.parse_args(argv)

    if args.gui:
        return gui_main()

    folder = ensure_folder(args.folder)
    checkin = parse_date(args.date)
    status_words = normalize_words(args.status_words)

    if args.update_db_target:
        update_single_campground_db(
            folder=folder,
            target=args.update_db_target,
            major=args.update_db_major,
            sheet_title=args.update_db_sheet_title,
            price=args.update_db_price,
            checkin=checkin,
            status_words=status_words,
            show_browser=not args.headless,
            wait_seconds=args.wait_seconds,
            progress=print,
        )
        return 0

    if args.auto_repair_rounds:
        run_collection_until_normal(
            folder=folder,
            checkin=checkin,
            status_words=status_words,
            total_mode=args.total_mode,
            show_browser=not args.headless,
            wait_seconds=args.wait_seconds,
            progress=print,
            max_rounds=args.auto_repair_rounds,
        )
    else:
        run_collection(
            folder=folder,
            checkin=checkin,
            status_words=status_words,
            total_mode=args.total_mode,
            db_update_mode=args.db_update_mode,
            show_browser=not args.headless,
            wait_seconds=args.wait_seconds,
            progress=print,
        )
    return 0


def gui_main() -> int:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class BookingApp(tk.Tk):
        def __init__(self) -> None:
            super().__init__()
            self.title(APP_TITLE)
            self.geometry("1180x760")
            self.minsize(1040, 680)
            self.stop_event = threading.Event()
            self.worker: threading.Thread | None = None

            self.folder_var = tk.StringVar(value=DEFAULT_FOLDER)
            self.date_var = tk.StringVar(value=date.today().strftime("%Y-%m-%d"))
            self.status_var = tk.StringVar(value=DEFAULT_STATUS_WORDS)
            self.total_mode_var = tk.StringVar(value="matched")
            self.db_update_mode_var = tk.StringVar(value="candidates")
            self.show_browser_var = tk.BooleanVar(value=True)
            self.wait_var = tk.IntVar(value=25)
            self.auto_batch_repair_var = tk.BooleanVar(value=False)
            self.auto_batch_rounds_var = tk.IntVar(value=5)
            self.new_target_var = tk.StringVar(value="")
            self.new_major_var = tk.StringVar(value="")
            self.new_sheet_title_var = tk.StringVar(value="")
            self.new_price_var = tk.StringVar(value="")
            self.rerun_after_db_update_var = tk.BooleanVar(value=False)
            self.summary_var = tk.StringVar(value="수집 전")
            self.db_status_var = tk.StringVar(value="")
            self.db_search_var = tk.StringVar(value="")
            self.edit_row_var = tk.StringVar(value="")
            self.edit_company_var = tk.StringVar(value="")
            self.edit_major_var = tk.StringVar(value="")
            self.edit_sheet_var = tk.StringVar(value="")
            self.edit_price_var = tk.StringVar(value="")
            self.db_entries_by_row: dict[int, DbEntry] = {}
            self.status_rows_by_iid: dict[str, dict[str, object]] = {}
            self.last_issue_rows: set[int] = set()

            self._build_ui()

        def _build_ui(self) -> None:
            outer = ttk.Frame(self, padding=14)
            outer.pack(fill="both", expand=True)

            title = ttk.Label(outer, text=APP_TITLE, font=("Malgun Gothic", 16, "bold"))
            title.pack(anchor="w", pady=(0, 10))

            form = ttk.Frame(outer)
            form.pack(fill="x")
            form.columnconfigure(1, weight=1)

            ttk.Label(form, text="프로그램 폴더").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
            ttk.Entry(form, textvariable=self.folder_var).grid(row=0, column=1, sticky="ew", pady=4)
            ttk.Button(form, text="찾기", command=self.browse_folder).grid(row=0, column=2, padx=(8, 0), pady=4)

            ttk.Label(form, text="체크인 날짜").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
            ttk.Entry(form, textvariable=self.date_var, width=16).grid(row=1, column=1, sticky="w", pady=4)

            ttk.Label(form, text="예약 완료 문구").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=4)
            ttk.Entry(form, textvariable=self.status_var).grid(row=2, column=1, columnspan=2, sticky="ew", pady=4)

            ttk.Label(form, text="총 객실수 기준").grid(row=3, column=0, sticky="w", padx=(0, 8), pady=4)
            total_mode = ttk.Combobox(
                form,
                textvariable=self.total_mode_var,
                state="readonly",
                width=28,
                values=[
                    "matched",
                    "db",
                    "max",
                ],
            )
            total_mode.grid(row=3, column=1, sticky="w", pady=4)
            ttk.Label(form, text="matched=기존 방식, db=DB 기준, max=둘 중 큰 값").grid(
                row=3, column=2, sticky="w", padx=(8, 0), pady=4
            )

            ttk.Label(form, text="DB 객실명 보정").grid(row=4, column=0, sticky="w", padx=(0, 8), pady=4)
            db_update_mode = ttk.Combobox(
                form,
                textvariable=self.db_update_mode_var,
                state="readonly",
                width=28,
                values=[
                    "candidates",
                    "auto-safe",
                    "auto-all",
                    "off",
                ],
            )
            db_update_mode.grid(row=4, column=1, sticky="w", pady=4)
            ttk.Label(form, text="auto는 단일 업체번호만 자동 반영, 혼합업체는 DB관리에서 분류").grid(
                row=4, column=2, sticky="w", padx=(8, 0), pady=4
            )

            options = ttk.Frame(outer)
            options.pack(fill="x", pady=(8, 6))
            ttk.Checkbutton(options, text="브라우저 화면 표시", variable=self.show_browser_var).pack(side="left")
            ttk.Label(options, text="대기시간").pack(side="left", padx=(18, 4))
            ttk.Spinbox(options, from_=5, to=90, width=6, textvariable=self.wait_var).pack(side="left")
            ttk.Label(options, text="초").pack(side="left", padx=(4, 0))
            ttk.Checkbutton(
                options,
                text="매칭실패 전체 일괄수정 후 정상까지 재수집",
                variable=self.auto_batch_repair_var,
            ).pack(side="left", padx=(18, 4))
            ttk.Label(options, text="최대").pack(side="left")
            ttk.Spinbox(options, from_=1, to=10, width=4, textvariable=self.auto_batch_rounds_var).pack(side="left", padx=(4, 2))
            ttk.Label(options, text="회").pack(side="left")

            db_frame = ttk.LabelFrame(outer, text="신규/특정 캠핑장 DB 추가 또는 업데이트", padding=10)
            db_frame.pack(fill="x", pady=(4, 10))
            db_frame.columnconfigure(1, weight=1)
            db_frame.columnconfigure(3, weight=1)

            ttk.Label(db_frame, text="네이버 링크/업체번호").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=3)
            ttk.Entry(db_frame, textvariable=self.new_target_var).grid(row=0, column=1, columnspan=3, sticky="ew", pady=3)
            ttk.Label(db_frame, text="대분류").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=3)
            ttk.Entry(db_frame, textvariable=self.new_major_var).grid(row=1, column=1, sticky="ew", pady=3)
            ttk.Label(db_frame, text="중분류/시트명").grid(row=1, column=2, sticky="w", padx=(12, 8), pady=3)
            ttk.Entry(db_frame, textvariable=self.new_sheet_title_var).grid(row=1, column=3, sticky="ew", pady=3)
            ttk.Label(db_frame, text="금액").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=3)
            ttk.Entry(db_frame, textvariable=self.new_price_var, width=18).grid(row=2, column=1, sticky="w", pady=3)
            self.db_update_btn = ttk.Button(db_frame, text="DB 추가/업데이트", command=self.start_single_db_update)
            self.db_update_btn.grid(row=0, column=4, rowspan=3, sticky="ns", padx=(10, 0), pady=3)
            ttk.Checkbutton(db_frame, text="DB 수정 후 수정한 항목만 재수집", variable=self.rerun_after_db_update_var).grid(
                row=2, column=2, columnspan=2, sticky="w", pady=(6, 0)
            )

            buttons = ttk.Frame(outer)
            buttons.pack(fill="x", pady=(4, 10))
            self.start_btn = ttk.Button(buttons, text="수집 시작", command=self.start)
            self.start_btn.pack(side="left")
            self.stop_btn = ttk.Button(buttons, text="중지", command=self.stop, state="disabled")
            self.stop_btn.pack(side="left", padx=(8, 0))
            self.issue_rerun_btn = ttk.Button(buttons, text="알림항목 재검색/반영", command=self.rerun_issue_rows)
            self.issue_rerun_btn.pack(side="left", padx=(8, 0))
            ttk.Button(buttons, text="로그 지우기", command=lambda: self.log_text.delete("1.0", "end")).pack(
                side="left", padx=(8, 0)
            )

            self.progress = ttk.Progressbar(outer, mode="indeterminate")
            self.progress.pack(fill="x", pady=(0, 8))

            ttk.Label(outer, textvariable=self.summary_var, font=("Malgun Gothic", 10, "bold")).pack(
                fill="x", pady=(0, 6)
            )

            self.notebook = ttk.Notebook(outer)
            self.notebook.pack(fill="both", expand=True)

            log_tab = ttk.Frame(self.notebook, padding=8)
            result_tab = ttk.Frame(self.notebook, padding=8)
            status_tab = ttk.Frame(self.notebook, padding=8)
            db_tab = ttk.Frame(self.notebook, padding=8)
            self.notebook.add(log_tab, text="수집 로그")
            self.notebook.add(result_tab, text="수집 결과")
            self.notebook.add(status_tab, text="현황")
            self.notebook.add(db_tab, text="객실DB 관리")
            self.status_tab = status_tab
            self.db_tab = db_tab

            self.log_text = tk.Text(log_tab, height=20, wrap="word", font=("Malgun Gothic", 10))
            self.log_text.pack(fill="both", expand=True)
            self._build_result_tab(result_tab)
            self._build_status_tab(status_tab)
            self._build_db_tab(db_tab)

            self.log("기존 폴더를 선택한 뒤 수집 시작을 누르세요.")
            self.log("필요 파일: 객실_DB.xlsx, 예약현황_Result.xlsx")
            self.log("매칭 실패가 있으면 백업 폴더에 DB 업데이트 후보 CSV가 저장됩니다.")
            self.refresh_db_tree()
            self.refresh_status_tree()

        def _build_result_tab(self, parent) -> None:
            parent.rowconfigure(0, weight=1)
            parent.columnconfigure(0, weight=1)
            columns = (
                "round",
                "row",
                "company",
                "major",
                "sheet",
                "price",
                "reserved",
                "matched",
                "missing",
                "status",
            )
            self.result_tree = ttk.Treeview(parent, columns=columns, show="headings", height=14)
            headings = {
                "round": "회차",
                "row": "DB행",
                "company": "고유인덱스",
                "major": "대분류",
                "sheet": "중분류",
                "price": "금액",
                "reserved": "예약/총",
                "matched": "매칭/DB",
                "missing": "누락",
                "status": "결과",
            }
            widths = {
                "round": 50,
                "row": 60,
                "company": 100,
                "major": 120,
                "sheet": 160,
                "price": 90,
                "reserved": 80,
                "matched": 80,
                "missing": 60,
                "status": 360,
            }
            for column in columns:
                self.result_tree.heading(column, text=headings[column])
                self.result_tree.column(column, width=widths[column], anchor="center" if column in {"round", "row", "reserved", "matched", "missing"} else "w")
            result_scroll = ttk.Scrollbar(parent, orient="vertical", command=self.result_tree.yview)
            self.result_tree.configure(yscrollcommand=result_scroll.set)
            self.result_tree.grid(row=0, column=0, sticky="nsew")
            result_scroll.grid(row=0, column=1, sticky="ns")
            self.result_tree.bind("<Double-1>", self.load_result_selection_into_db_editor)

        def _build_status_tab(self, parent) -> None:
            parent.rowconfigure(1, weight=1)
            parent.columnconfigure(0, weight=2)
            parent.columnconfigure(1, weight=3)

            toolbar = ttk.Frame(parent)
            toolbar.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 8))
            ttk.Button(toolbar, text="현황 새로고침", command=self.refresh_status_tree).pack(side="left")
            ttk.Label(toolbar, text="DB없음은 결과파일에는 남아 있지만 현재 객실DB에는 없는 옛 시트입니다.").pack(
                side="left", padx=(10, 0)
            )

            columns = ("state", "row", "company", "major", "sheet", "price", "latest", "reserved", "rate", "revenue")
            self.status_tree = ttk.Treeview(parent, columns=columns, show="headings", height=16)
            headings = {
                "state": "상태",
                "row": "DB행",
                "company": "고유인덱스",
                "major": "대분류",
                "sheet": "중분류",
                "price": "금액",
                "latest": "최근일자",
                "reserved": "예약/총",
                "rate": "가동률",
                "revenue": "추정매출",
            }
            widths = {
                "state": 70,
                "row": 55,
                "company": 95,
                "major": 120,
                "sheet": 160,
                "price": 85,
                "latest": 90,
                "reserved": 70,
                "rate": 70,
                "revenue": 95,
            }
            for column in columns:
                self.status_tree.heading(column, text=headings[column])
                self.status_tree.column(column, width=widths[column], anchor="center" if column in {"state", "row", "latest", "reserved", "rate"} else "w")
            status_scroll = ttk.Scrollbar(parent, orient="vertical", command=self.status_tree.yview)
            self.status_tree.configure(yscrollcommand=status_scroll.set)
            self.status_tree.grid(row=1, column=0, sticky="nsew", padx=(0, 8))
            status_scroll.grid(row=1, column=0, sticky="nse", padx=(0, 8))
            self.status_tree.bind("<<TreeviewSelect>>", self.load_selected_status_history)
            self.status_tree.bind("<Double-1>", self.load_status_selection_into_db_editor)

            history_frame = ttk.LabelFrame(parent, text="날짜별 현황", padding=8)
            history_frame.grid(row=1, column=1, sticky="nsew")
            history_frame.rowconfigure(0, weight=1)
            history_frame.columnconfigure(0, weight=1)
            history_columns = ("date", "reserved", "total", "rate", "price", "revenue")
            self.history_tree = ttk.Treeview(history_frame, columns=history_columns, show="headings", height=16)
            history_headings = {
                "date": "일자",
                "reserved": "예약",
                "total": "총객실",
                "rate": "가동률",
                "price": "금액",
                "revenue": "추정매출",
            }
            history_widths = {"date": 100, "reserved": 60, "total": 60, "rate": 70, "price": 90, "revenue": 110}
            for column in history_columns:
                self.history_tree.heading(column, text=history_headings[column])
                self.history_tree.column(column, width=history_widths[column], anchor="center")
            history_scroll = ttk.Scrollbar(history_frame, orient="vertical", command=self.history_tree.yview)
            self.history_tree.configure(yscrollcommand=history_scroll.set)
            self.history_tree.grid(row=0, column=0, sticky="nsew")
            history_scroll.grid(row=0, column=1, sticky="ns")

        def _build_db_tab(self, parent) -> None:
            parent.rowconfigure(1, weight=1)
            parent.columnconfigure(0, weight=1)
            parent.columnconfigure(1, weight=1)

            toolbar = ttk.Frame(parent)
            toolbar.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 8))
            ttk.Label(toolbar, text="검색").pack(side="left")
            ttk.Entry(toolbar, textvariable=self.db_search_var, width=28).pack(side="left", padx=(6, 4))
            ttk.Button(toolbar, text="새로고침", command=self.refresh_db_tree).pack(side="left", padx=(4, 0))
            ttk.Button(toolbar, text="복사하여 추가", command=self.copy_selected_db_entry).pack(side="left", padx=(4, 0))
            ttk.Button(toolbar, text="선택 삭제", command=self.delete_selected_db_entry).pack(side="left", padx=(4, 0))
            ttk.Label(toolbar, textvariable=self.db_status_var).pack(side="right")

            columns = ("row", "company", "major", "sheet", "price", "rooms")
            self.db_tree = ttk.Treeview(parent, columns=columns, show="headings", height=16)
            headings = {
                "row": "DB행",
                "company": "고유인덱스",
                "major": "대분류",
                "sheet": "중분류/시트명",
                "price": "금액",
                "rooms": "객실수",
            }
            widths = {"row": 55, "company": 110, "major": 130, "sheet": 180, "price": 90, "rooms": 65}
            for column in columns:
                self.db_tree.heading(column, text=headings[column])
                self.db_tree.column(column, width=widths[column], anchor="center" if column in {"row", "rooms"} else "w")
            db_scroll = ttk.Scrollbar(parent, orient="vertical", command=self.db_tree.yview)
            self.db_tree.configure(yscrollcommand=db_scroll.set)
            self.db_tree.grid(row=1, column=0, sticky="nsew", padx=(0, 8))
            db_scroll.grid(row=1, column=0, sticky="nse", padx=(0, 8))
            self.db_tree.bind("<<TreeviewSelect>>", self.on_db_select)

            editor = ttk.LabelFrame(parent, text="선택한 DB 행 수정", padding=10)
            editor.grid(row=1, column=1, sticky="nsew")
            editor.columnconfigure(1, weight=1)
            editor.rowconfigure(5, weight=1)

            ttk.Label(editor, text="DB행").grid(row=0, column=0, sticky="w", pady=3)
            ttk.Entry(editor, textvariable=self.edit_row_var, width=12, state="readonly").grid(row=0, column=1, sticky="w", pady=3)
            ttk.Label(editor, text="고유인덱스").grid(row=1, column=0, sticky="w", pady=3)
            ttk.Entry(editor, textvariable=self.edit_company_var).grid(row=1, column=1, sticky="ew", pady=3)
            ttk.Label(editor, text="대분류").grid(row=2, column=0, sticky="w", pady=3)
            ttk.Entry(editor, textvariable=self.edit_major_var).grid(row=2, column=1, sticky="ew", pady=3)
            ttk.Label(editor, text="중분류/시트명").grid(row=3, column=0, sticky="w", pady=3)
            ttk.Entry(editor, textvariable=self.edit_sheet_var).grid(row=3, column=1, sticky="ew", pady=3)
            ttk.Label(editor, text="금액").grid(row=4, column=0, sticky="w", pady=3)
            ttk.Entry(editor, textvariable=self.edit_price_var, width=18).grid(row=4, column=1, sticky="w", pady=3)
            ttk.Label(editor, text="세분류/객실명").grid(row=5, column=0, sticky="nw", pady=3)
            self.room_text = tk.Text(editor, height=14, wrap="word", font=("Malgun Gothic", 10))
            self.room_text.grid(row=5, column=1, sticky="nsew", pady=3)
            room_scroll = ttk.Scrollbar(editor, orient="vertical", command=self.room_text.yview)
            self.room_text.configure(yscrollcommand=room_scroll.set)
            room_scroll.grid(row=5, column=2, sticky="ns", pady=3)

            edit_buttons = ttk.Frame(editor)
            edit_buttons.grid(row=6, column=1, sticky="ew", pady=(8, 0))
            ttk.Button(edit_buttons, text="저장", command=self.save_selected_db_entry).pack(side="left")
            ttk.Button(
                edit_buttons,
                text="저장 후 이 항목 재수집",
                command=lambda: self.save_selected_db_entry(rerun=True),
            ).pack(side="left", padx=(6, 0))
            ttk.Button(edit_buttons, text="네이버 객실 불러오기", command=self.load_rooms_from_naver_to_editor).pack(
                side="left", padx=(6, 0)
            )
            ttk.Button(edit_buttons, text="같은 업체 자동분리", command=self.auto_split_company_rooms).pack(
                side="left", padx=(6, 0)
            )

        def browse_folder(self) -> None:
            folder = filedialog.askdirectory(initialdir=self.folder_var.get() or DEFAULT_FOLDER)
            if folder:
                self.folder_var.set(folder)
                self.refresh_db_tree()
                self.refresh_status_tree()

        def log(self, message: str) -> None:
            self.log_text.insert("end", message + "\n")
            self.log_text.see("end")
            self.update_idletasks()

        def money_display(self, value: object) -> str:
            return format_price_text(value)

        def rate_display(self, value: object) -> str:
            if value in (None, ""):
                return ""
            try:
                return f"{float(value) * 100:.1f}%"
            except (TypeError, ValueError):
                return str(value)

        def clear_result_table(self) -> None:
            for item in self.result_tree.get_children():
                self.result_tree.delete(item)
            self.summary_var.set("수집 중")

        def add_result_row(self, round_no: int, result: EntryResult) -> None:
            status_label = "정상" if result.status == "정상" else "수정필요"
            self.result_tree.insert(
                "",
                "end",
                values=(
                    round_no,
                    result.entry.excel_row,
                    result.entry.company_id,
                    result.entry.major,
                    result.entry.sheet_title,
                    result.entry.price,
                    f"{result.reserved_count}/{result.saved_total}",
                    f"{result.matched_count}/{result.db_room_count}",
                    result.missing_count,
                    f"{status_label} - {result.status}",
                ),
            )

        def refresh_status_tree(self) -> None:
            try:
                folder = ensure_folder(self.folder_var.get())
                rows = read_status_rows(folder)
            except Exception as exc:
                for item in self.status_tree.get_children():
                    self.status_tree.delete(item)
                for item in self.history_tree.get_children():
                    self.history_tree.delete(item)
                self.status_rows_by_iid = {}
                self.summary_var.set(f"현황 읽기 실패: {exc}")
                return

            self.status_rows_by_iid = {}
            for item in self.status_tree.get_children():
                self.status_tree.delete(item)
            for item in self.history_tree.get_children():
                self.history_tree.delete(item)

            for index, row in enumerate(rows, start=1):
                entry = row.get("entry")
                latest = row.get("latest") or {}
                sheet = str(row.get("sheet") or "")
                reserved = safe_int(latest.get("reserved")) if latest else 0
                total = safe_int(latest.get("total")) if latest else (len(entry.rooms) if entry else 0)
                rate = calc_rate(reserved, total)
                price = entry.price if entry else latest.get("price", "")
                revenue = reserved * parse_price_amount(price) if parse_price_amount(price) is not None else latest.get("revenue", "")
                iid = f"status{index}"
                self.status_rows_by_iid[iid] = row
                self.status_tree.insert(
                    "",
                    "end",
                    iid=iid,
                    values=(
                        row.get("state", ""),
                        entry.excel_row if entry else "",
                        entry.company_id if entry else "",
                        entry.major if entry else "",
                        sheet,
                        self.money_display(price),
                        latest.get("date_text", "") if latest else "",
                        f"{reserved}/{total}" if latest else f"-/{total}" if total else "",
                        self.rate_display(rate),
                        self.money_display(revenue),
                    ),
                )

        def load_selected_status_history(self, _event=None) -> None:
            selected = self.status_tree.selection()
            if not selected:
                return
            row = self.status_rows_by_iid.get(selected[0])
            if not row:
                return
            sheet = str(row.get("sheet") or "")
            try:
                folder = ensure_folder(self.folder_var.get())
                history = read_result_history(folder, sheet)
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))
                return

            for item in self.history_tree.get_children():
                self.history_tree.delete(item)
            for item in history:
                self.history_tree.insert(
                    "",
                    "end",
                    values=(
                        item.get("date_text", ""),
                        item.get("reserved", ""),
                        item.get("total", ""),
                        self.rate_display(item.get("rate")),
                        self.money_display(item.get("price")),
                        self.money_display(item.get("revenue")),
                    ),
                )

        def load_status_selection_into_db_editor(self, _event=None) -> None:
            selected = self.status_tree.selection()
            if not selected:
                return
            row = self.status_rows_by_iid.get(selected[0])
            entry = row.get("entry") if row else None
            if entry:
                self.load_entry_to_editor(entry)

        def refresh_db_tree(self) -> None:
            try:
                folder = ensure_folder(self.folder_var.get())
                entries = load_db(folder / DB_FILENAME)
            except Exception as exc:
                self.db_status_var.set(f"DB 읽기 실패: {exc}")
                return

            self.db_entries_by_row = {entry.excel_row: entry for entry in entries}
            search_key = normalize_name(self.db_search_var.get()).lower()
            for item in self.db_tree.get_children():
                self.db_tree.delete(item)
            for entry in entries:
                haystack = normalize_name(
                    f"{entry.company_id} {entry.major} {entry.sheet_title} {entry.price} {' '.join(entry.rooms[:8])}"
                ).lower()
                if search_key and search_key not in haystack:
                    continue
                self.db_tree.insert(
                    "",
                    "end",
                    iid=str(entry.excel_row),
                    values=(
                        entry.excel_row,
                        entry.company_id,
                        entry.major,
                        entry.sheet_title,
                        entry.price,
                        len(entry.rooms),
                    ),
                )
            self.db_status_var.set(f"DB {len(entries)}개 행")

        def select_db_row(self, row_text: str) -> None:
            if row_text in self.db_tree.get_children():
                self.db_tree.selection_set(row_text)
                self.db_tree.see(row_text)
                self.on_db_select()

        def on_db_select(self, _event=None) -> None:
            selected = self.db_tree.selection()
            if not selected:
                return
            entry = self.db_entries_by_row.get(int(selected[0]))
            if not entry:
                return
            self.load_entry_to_editor(entry)

        def load_entry_to_editor(self, entry: DbEntry) -> None:
            self.edit_row_var.set(str(entry.excel_row))
            self.edit_company_var.set(entry.company_id)
            self.edit_major_var.set(entry.major)
            self.edit_sheet_var.set(entry.sheet_title)
            self.edit_price_var.set(entry.price)
            self.room_text.delete("1.0", "end")
            self.room_text.insert("1.0", "\n".join(entry.rooms))
            self.notebook.select(self.db_tab)

        def load_result_selection_into_db_editor(self, _event=None) -> None:
            selected = self.result_tree.selection()
            if not selected:
                return
            values = self.result_tree.item(selected[0], "values")
            if len(values) < 2:
                return
            try:
                row_idx = int(values[1])
            except ValueError:
                return
            entry = self.db_entries_by_row.get(row_idx)
            if not entry:
                self.refresh_db_tree()
                entry = self.db_entries_by_row.get(row_idx)
            if entry:
                self.load_entry_to_editor(entry)

        def new_db_entry(self) -> None:
            self.edit_row_var.set("")
            self.edit_company_var.set("")
            self.edit_major_var.set("")
            self.edit_sheet_var.set("")
            self.edit_price_var.set("")
            self.room_text.delete("1.0", "end")
            self.notebook.select(self.db_tab)

        def copy_selected_db_entry(self) -> None:
            selected = self.db_tree.selection()
            if not selected:
                messagebox.showwarning(APP_TITLE, "복사할 DB 행을 선택해 주세요.")
                return
            entry = self.db_entries_by_row.get(int(selected[0]))
            if not entry:
                return
            self.edit_row_var.set("")
            self.edit_company_var.set(entry.company_id)
            self.edit_major_var.set(entry.major)
            self.edit_sheet_var.set(f"{entry.sheet_title}_복사")
            self.edit_price_var.set(entry.price)
            self.room_text.delete("1.0", "end")
            self.room_text.insert("1.0", "\n".join(entry.rooms))
            self.notebook.select(self.db_tab)
            self.log(f"[{log_time()}] {entry.sheet_title} 행을 복사했습니다. 중분류/객실명을 수정한 뒤 저장하세요.")

        def editor_rooms(self) -> list[str]:
            return split_room_text(self.room_text.get("1.0", "end"))

        def save_selected_db_entry(self, rerun: bool = False) -> None:
            if rerun and self.worker and self.worker.is_alive():
                return
            try:
                folder = ensure_folder(self.folder_var.get())
                db_path = folder / DB_FILENAME
                row_text = self.edit_row_var.get().strip()
                row_idx = int(row_text) if row_text else None
                if rerun:
                    checkin = parse_date(self.date_var.get())
                    status_words = normalize_words(self.status_var.get())
                    total_mode = self.total_mode_var.get()
                    db_update_mode = self.db_update_mode_var.get()
                    show_browser = self.show_browser_var.get()
                    wait_seconds = int(self.wait_var.get())
                    if not status_words:
                        raise ValueError("예약 완료 문구가 비어 있습니다.")
                backup_db_file(db_path)
                saved_row = save_db_entry(
                    db_path=db_path,
                    row_idx=row_idx,
                    company_id=self.edit_company_var.get(),
                    major=self.edit_major_var.get(),
                    sheet_title=self.edit_sheet_var.get(),
                    price=self.edit_price_var.get(),
                    rooms=self.editor_rooms(),
                )
                self.log(f"[{log_time()}] DB 저장 완료: {saved_row}행")
                self.refresh_db_tree()
                self.refresh_status_tree()
                entry = self.db_entries_by_row.get(saved_row)
                if entry:
                    self.load_entry_to_editor(entry)
                if rerun:
                    self.stop_event.clear()
                    self.start_btn.configure(state="disabled")
                    self.db_update_btn.configure(state="disabled")
                    self.stop_btn.configure(state="normal")
                    self.progress.start(12)
                    self.clear_result_table()
                    self.notebook.select(1)
                    self.log(f"[{log_time()}] 저장한 DB {saved_row}행만 재수집을 시작합니다.")

                    def worker() -> None:
                        try:
                            results = run_collection(
                                folder=folder,
                                checkin=checkin,
                                status_words=status_words,
                                total_mode=total_mode,
                                db_update_mode=db_update_mode,
                                show_browser=show_browser,
                                wait_seconds=wait_seconds,
                                progress=lambda msg: self.after(0, self.log, msg),
                                stop_event=self.stop_event,
                                result_callback=lambda result: self.after(0, self.add_result_row, 1, result),
                                target_excel_rows={saved_row},
                            )
                            normal = sum(1 for item in results if item.status == "정상")
                            issue_count = len(results) - normal
                            self.after(
                                0,
                                self.summary_var.set,
                                f"부분 재수집 완료: 수정항목 {len(results)}개 중 정상 {normal}개 / 수정필요 {issue_count}개",
                            )
                            self.after(0, self.log, f"[{log_time()}] 부분 재수집 완료: {normal}/{len(results)}개 정상")
                            self.after(0, self.refresh_status_tree)
                            self.after(0, self.warn_if_mismatch, results)
                        except Exception as exc:
                            self.after(0, self.summary_var.set, f"오류: {exc}")
                            self.after(0, self.log, f"[{log_time()}] 오류: {exc}")
                            self.after(0, messagebox.showerror, APP_TITLE, str(exc))
                        finally:
                            self.after(0, self.finish)

                    self.worker = threading.Thread(target=worker, daemon=True)
                    self.worker.start()
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))

        def delete_selected_db_entry(self) -> None:
            selected = self.db_tree.selection()
            if not selected:
                messagebox.showwarning(APP_TITLE, "삭제할 DB 행을 선택해 주세요.")
                return
            row_idx = int(selected[0])
            entry = self.db_entries_by_row.get(row_idx)
            if not entry:
                return
            if not messagebox.askyesno(APP_TITLE, f"{entry.sheet_title} DB 행을 삭제할까요?"):
                return
            try:
                folder = ensure_folder(self.folder_var.get())
                db_path = folder / DB_FILENAME
                backup_db_file(db_path)
                delete_db_entry(db_path, row_idx)
                self.log(f"[{log_time()}] DB 삭제 완료: {row_idx}행 {entry.sheet_title}")
                self.new_db_entry()
                self.refresh_db_tree()
                self.refresh_status_tree()
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))

        def load_rooms_from_naver_to_editor(self) -> None:
            if self.worker and self.worker.is_alive():
                return
            try:
                company_id = parse_company_id(self.edit_company_var.get())
                checkin = parse_date(self.date_var.get())
                status_words = normalize_words(self.status_var.get())
                show_browser = self.show_browser_var.get()
                wait_seconds = int(self.wait_var.get())
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))
                return

            self.progress.start(12)
            self.start_btn.configure(state="disabled")
            self.db_update_btn.configure(state="disabled")
            self.log(f"[{log_time()}] 네이버 객실 불러오기 시작: {company_id}")

            def worker() -> None:
                try:
                    with BrowserCollector(checkin, status_words, show_browser, wait_seconds, lambda msg: self.after(0, self.log, msg)) as collector:
                        _url, snapshot = collector.get_snapshot(company_id)
                        rooms = extract_detected_rooms(snapshot, status_words)
                    if not rooms:
                        raise ValueError("네이버 화면에서 객실명을 찾지 못했습니다.")
                    self.after(0, self.room_text.delete, "1.0", "end")
                    self.after(0, self.room_text.insert, "1.0", "\n".join(rooms))
                    self.after(0, self.log, f"[{log_time()}] 네이버 객실 {len(rooms)}개 불러옴. 확인 후 [저장]을 누르세요.")
                except Exception as exc:
                    self.after(0, self.summary_var.set, f"오류: {exc}")
                    self.after(0, self.log, f"[{log_time()}] 오류: {exc}")
                    self.after(0, messagebox.showerror, APP_TITLE, str(exc))
                finally:
                    self.after(0, self.finish)

            self.worker = threading.Thread(target=worker, daemon=True)
            self.worker.start()

        def auto_split_company_rooms(self) -> None:
            if self.worker and self.worker.is_alive():
                return
            try:
                folder = ensure_folder(self.folder_var.get())
                db_path = folder / DB_FILENAME
                company_id = parse_company_id(self.edit_company_var.get())
                checkin = parse_date(self.date_var.get())
                status_words = normalize_words(self.status_var.get())
                show_browser = self.show_browser_var.get()
                wait_seconds = int(self.wait_var.get())
                rerun_after_update = bool(self.rerun_after_db_update_var.get())
                total_mode = self.total_mode_var.get()
                db_update_mode = self.db_update_mode_var.get()
                entries = [entry for entry in load_db(db_path) if entry.company_id == company_id]
                if len(entries) < 2:
                    raise ValueError("같은 고유인덱스번호 DB 행이 2개 이상일 때만 자동분리할 수 있습니다.")
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))
                return

            self.progress.start(12)
            self.start_btn.configure(state="disabled")
            self.db_update_btn.configure(state="disabled")
            self.log(f"[{log_time()}] 같은 업체 자동분리 시작: {company_id}")

            def worker() -> None:
                try:
                    with BrowserCollector(checkin, status_words, show_browser, wait_seconds, lambda msg: self.after(0, self.log, msg)) as collector:
                        _url, snapshot = collector.get_snapshot(company_id)
                        rooms = extract_detected_rooms(snapshot, status_words)
                    updates, note = split_rooms_by_product_type(entries, rooms)
                    if not updates:
                        raise ValueError(note)
                    db_backup_path = backup_db_file(db_path)
                    if db_backup_path:
                        self.after(0, self.log, f"[{log_time()}] DB 파일 백업 완료: {db_backup_path.name}")
                    updated_count = write_db_rooms(db_path, updates)
                    self.after(0, self.log, f"[{log_time()}] 같은 업체 자동분리 완료: {updated_count}개 행, {note}")
                    self.after(0, self.refresh_db_tree)
                    self.after(0, self.refresh_status_tree)
                    selected_row = self.edit_row_var.get().strip()
                    if selected_row:
                        self.after(0, self.select_db_row, selected_row)
                    if rerun_after_update:
                        target_rows = set(updates)
                        self.after(0, self.log, f"[{log_time()}] 자동분리로 수정된 {len(target_rows)}개 행만 재수집합니다.")
                        self.after(0, self.clear_result_table)
                        self.after(0, self.notebook.select, 1)
                        results = run_collection(
                            folder=folder,
                            checkin=checkin,
                            status_words=status_words,
                            total_mode=total_mode,
                            db_update_mode=db_update_mode,
                            show_browser=show_browser,
                            wait_seconds=wait_seconds,
                            progress=lambda msg: self.after(0, self.log, msg),
                            stop_event=self.stop_event,
                            result_callback=lambda result: self.after(0, self.add_result_row, 1, result),
                            target_excel_rows=target_rows,
                        )
                        normal = sum(1 for item in results if item.status == "정상")
                        issue_count = len(results) - normal
                        self.after(
                            0,
                            self.summary_var.set,
                            f"부분 재수집 완료: 수정항목 {len(results)}개 중 정상 {normal}개 / 수정필요 {issue_count}개",
                        )
                        self.after(0, self.log, f"[{log_time()}] 부분 재수집 완료: {normal}/{len(results)}개 정상")
                        self.after(0, self.refresh_status_tree)
                        self.after(0, self.warn_if_mismatch, results)
                except Exception as exc:
                    self.after(0, self.summary_var.set, f"오류: {exc}")
                    self.after(0, self.log, f"[{log_time()}] 오류: {exc}")
                    self.after(0, messagebox.showerror, APP_TITLE, str(exc))
                finally:
                    self.after(0, self.finish)

            self.worker = threading.Thread(target=worker, daemon=True)
            self.worker.start()

        def warn_if_mismatch(self, results: list[EntryResult]) -> None:
            issues = [
                item
                for item in results
                if result_needs_db_repair(item)
            ]
            if not issues:
                self.last_issue_rows.clear()
                return
            self.last_issue_rows = {item.entry.excel_row for item in issues}

            first = issues[0]
            self.new_target_var.set(first.entry.company_id)
            self.new_major_var.set(first.entry.major)
            self.new_sheet_title_var.set(first.entry.sheet_title)
            self.new_price_var.set(first.entry.price)

            lines = [
                f"객실 정보가 맞지 않는 항목이 {len(issues)}개 있습니다.",
                "",
                "첫 문제 항목을 아래 DB 수정 입력칸에 자동으로 넣었습니다.",
                "확인 후 [DB 추가/업데이트]를 누르면 네이버에서 객실명을 다시 찾아 DB를 수정합니다.",
                "[알림항목 재검색/반영]을 누르면 아래 목록의 문제 항목만 다시 처리합니다.",
                "수집 결과 표에서 원하는 행을 선택한 뒤 누르면 선택 항목만 처리합니다.",
                "",
            ]
            for item in issues[:8]:
                lines.append(
                    f"- {item.entry.excel_row}행 {item.entry.sheet_title}: "
                    f"매칭 {item.matched_count}/DB {item.db_room_count}, 누락 {item.missing_count}"
                )
            if len(issues) > 8:
                lines.append(f"- 외 {len(issues) - 8}개")

            messagebox.showwarning(APP_TITLE, "\n".join(lines))

        def selected_or_last_issue_rows(self) -> set[int]:
            rows: set[int] = set()
            for item_id in self.result_tree.selection():
                values = self.result_tree.item(item_id, "values")
                if len(values) < 2:
                    continue
                try:
                    rows.add(int(values[1]))
                except (TypeError, ValueError):
                    pass
            return rows or set(self.last_issue_rows)

        def rerun_issue_rows(self) -> None:
            if self.worker and self.worker.is_alive():
                return
            target_rows = self.selected_or_last_issue_rows()
            if not target_rows:
                messagebox.showinfo(APP_TITLE, "재검색할 문제 항목이 없습니다. 먼저 수집을 실행하거나 수집 결과에서 행을 선택해 주세요.")
                return
            try:
                folder = ensure_folder(self.folder_var.get())
                checkin = parse_date(self.date_var.get())
                status_words = normalize_words(self.status_var.get())
                total_mode = self.total_mode_var.get()
                show_browser = self.show_browser_var.get()
                wait_seconds = int(self.wait_var.get())
                max_rounds = int(self.auto_batch_rounds_var.get())
                if not status_words:
                    raise ValueError("예약 완료 문구가 비어 있습니다.")
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))
                return

            self.stop_event.clear()
            self.start_btn.configure(state="disabled")
            self.db_update_btn.configure(state="disabled")
            self.stop_btn.configure(state="normal")
            self.progress.start(12)
            self.clear_result_table()
            self.notebook.select(1)
            self.log(f"[{log_time()}] 알림/선택 문제 항목 {len(target_rows)}개만 재검색/반영을 시작합니다.")

            def worker() -> None:
                try:
                    results = run_collection_until_normal(
                        folder=folder,
                        checkin=checkin,
                        status_words=status_words,
                        total_mode=total_mode,
                        show_browser=show_browser,
                        wait_seconds=wait_seconds,
                        progress=lambda msg: self.after(0, self.log, msg),
                        stop_event=self.stop_event,
                        max_rounds=max_rounds,
                        result_callback=lambda round_no, result: self.after(0, self.add_result_row, round_no, result),
                        initial_target_rows=target_rows,
                    )
                    normal = sum(1 for item in results if item.status == "정상")
                    issue_count = len(results) - normal
                    self.after(
                        0,
                        self.summary_var.set,
                        f"문제 항목 재검색 완료: 대상 {len(results)}개 중 정상 {normal}개 / 수정필요 {issue_count}개",
                    )
                    self.after(0, self.log, f"[{log_time()}] 문제 항목 재검색 완료: {normal}/{len(results)}개 정상")
                    self.after(0, self.refresh_db_tree)
                    self.after(0, self.refresh_status_tree)
                    self.after(0, self.warn_if_mismatch, results)
                except Exception as exc:
                    self.after(0, self.summary_var.set, f"오류: {exc}")
                    self.after(0, self.log, f"[{log_time()}] 오류: {exc}")
                    self.after(0, messagebox.showerror, APP_TITLE, str(exc))
                finally:
                    self.after(0, self.finish)

            self.worker = threading.Thread(target=worker, daemon=True)
            self.worker.start()

        def start(self) -> None:
            if self.worker and self.worker.is_alive():
                return
            try:
                folder = ensure_folder(self.folder_var.get())
                checkin = parse_date(self.date_var.get())
                status_words = normalize_words(self.status_var.get())
                auto_batch_repair = bool(self.auto_batch_repair_var.get())
                auto_batch_rounds = int(self.auto_batch_rounds_var.get())
                if not status_words:
                    raise ValueError("예약 완료 문구가 비어 있습니다.")
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))
                return

            self.stop_event.clear()
            self.start_btn.configure(state="disabled")
            self.db_update_btn.configure(state="disabled")
            self.stop_btn.configure(state="normal")
            self.progress.start(12)
            self.clear_result_table()
            self.notebook.select(1)
            self.log(f"[{log_time()}] 수집 시작")

            def worker() -> None:
                try:
                    if auto_batch_repair:
                        results = run_collection_until_normal(
                            folder=folder,
                            checkin=checkin,
                            status_words=status_words,
                            total_mode=self.total_mode_var.get(),
                            show_browser=self.show_browser_var.get(),
                            wait_seconds=int(self.wait_var.get()),
                            progress=lambda msg: self.after(0, self.log, msg),
                            stop_event=self.stop_event,
                            max_rounds=auto_batch_rounds,
                            result_callback=lambda round_no, result: self.after(0, self.add_result_row, round_no, result),
                        )
                    else:
                        results = run_collection(
                            folder=folder,
                            checkin=checkin,
                            status_words=status_words,
                            total_mode=self.total_mode_var.get(),
                            db_update_mode=self.db_update_mode_var.get(),
                            show_browser=self.show_browser_var.get(),
                            wait_seconds=int(self.wait_var.get()),
                            progress=lambda msg: self.after(0, self.log, msg),
                            stop_event=self.stop_event,
                            result_callback=lambda result: self.after(0, self.add_result_row, 1, result),
                        )
                    normal = sum(1 for item in results if item.status == "정상")
                    issue_count = len(results) - normal
                    self.after(0, self.summary_var.set, f"완료: 전체 {len(results)}개 중 정상 {normal}개 / 수정필요 {issue_count}개")
                    self.after(0, self.log, f"[{log_time()}] 완료: {normal}/{len(results)}개 정상")
                    self.after(0, self.refresh_db_tree)
                    self.after(0, self.refresh_status_tree)
                    if not auto_batch_repair:
                        self.after(0, self.warn_if_mismatch, results)
                except Exception as exc:
                    self.after(0, self.summary_var.set, f"오류: {exc}")
                    self.after(0, self.log, f"[{log_time()}] 오류: {exc}")
                    self.after(0, messagebox.showerror, APP_TITLE, str(exc))
                finally:
                    self.after(0, self.finish)

            self.worker = threading.Thread(target=worker, daemon=True)
            self.worker.start()

        def start_single_db_update(self) -> None:
            if self.worker and self.worker.is_alive():
                return
            try:
                folder = ensure_folder(self.folder_var.get())
                checkin = parse_date(self.date_var.get())
                status_words = normalize_words(self.status_var.get())
                target = self.new_target_var.get().strip()
                major = self.new_major_var.get().strip()
                sheet_title = self.new_sheet_title_var.get().strip()
                price = self.new_price_var.get().strip()
                rerun_after_update = bool(self.rerun_after_db_update_var.get())
                total_mode = self.total_mode_var.get()
                db_update_mode = self.db_update_mode_var.get()
                show_browser = self.show_browser_var.get()
                wait_seconds = int(self.wait_var.get())
                if not target:
                    raise ValueError("네이버 링크 또는 업체 고유번호를 입력해 주세요.")
                if not major and not sheet_title:
                    raise ValueError("대분류 또는 중분류/시트명 중 하나는 입력해 주세요.")
            except Exception as exc:
                messagebox.showerror(APP_TITLE, str(exc))
                return

            self.stop_event.clear()
            self.start_btn.configure(state="disabled")
            self.db_update_btn.configure(state="disabled")
            self.stop_btn.configure(state="disabled")
            self.progress.start(12)
            self.log(f"[{log_time()}] 특정 캠핑장 DB 추가/업데이트 시작")

            def worker() -> None:
                try:
                    action, row, room_count = update_single_campground_db(
                        folder=folder,
                        target=target,
                        major=major,
                        sheet_title=sheet_title,
                        price=price,
                        checkin=checkin,
                        status_words=status_words,
                        show_browser=show_browser,
                        wait_seconds=wait_seconds,
                        progress=lambda msg: self.after(0, self.log, msg),
                    )
                    action_text = "추가" if action == "added" else "업데이트"
                    self.after(0, self.log, f"[{log_time()}] DB {action_text} 완료: {row}행, 객실 {room_count}개")
                    self.after(0, self.refresh_db_tree)
                    self.after(0, self.refresh_status_tree)
                    if rerun_after_update:
                        self.after(0, self.log, f"[{log_time()}] DB 수정 완료 후 {row}행만 재수집을 시작합니다.")
                        self.after(0, self.clear_result_table)
                        self.after(0, self.notebook.select, 1)
                        results = run_collection(
                            folder=folder,
                            checkin=checkin,
                            status_words=status_words,
                            total_mode=total_mode,
                            db_update_mode=db_update_mode,
                            show_browser=show_browser,
                            wait_seconds=wait_seconds,
                            progress=lambda msg: self.after(0, self.log, msg),
                            stop_event=self.stop_event,
                            result_callback=lambda result: self.after(0, self.add_result_row, 1, result),
                            target_excel_rows={row},
                        )
                        normal = sum(1 for item in results if item.status == "정상")
                        issue_count = len(results) - normal
                        self.after(0, self.summary_var.set, f"부분 재수집 완료: 수정항목 {len(results)}개 중 정상 {normal}개 / 수정필요 {issue_count}개")
                        self.after(0, self.log, f"[{log_time()}] 부분 재수집 완료: {normal}/{len(results)}개 정상")
                        self.after(0, self.refresh_status_tree)
                        self.after(0, self.warn_if_mismatch, results)
                except Exception as exc:
                    self.after(0, self.log, f"[{log_time()}] 오류: {exc}")
                    self.after(0, messagebox.showerror, APP_TITLE, str(exc))
                finally:
                    self.after(0, self.finish)

            self.worker = threading.Thread(target=worker, daemon=True)
            self.worker.start()

        def stop(self) -> None:
            self.stop_event.set()
            self.log(f"[{log_time()}] 중지 요청")

        def finish(self) -> None:
            self.progress.stop()
            self.start_btn.configure(state="normal")
            self.db_update_btn.configure(state="normal")
            self.stop_btn.configure(state="disabled")

    app = BookingApp()
    app.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(cli_main())
